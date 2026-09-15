#!/usr/bin/env python3
"""Generate EPMS role-based training decks (.pptx).

Produces three decks in this directory:
  - EPMS_Training_Requester.pptx
  - EPMS_Training_Procurement_Officer.pptx
  - EPMS_Training_AP_Clerk.pptx

Usage:  python generate_training_ppt.py
Requires: python-pptx
"""
from __future__ import annotations

import os

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

OUT_DIR = os.path.dirname(os.path.abspath(__file__))

# ── UniOps design tokens (packages/shell/tokens.css) ────────────────────────
TEAL_900 = RGBColor.from_string("053838")
TEAL_800 = RGBColor.from_string("064A4A")
TEAL_700 = RGBColor.from_string("085E5E")   # primary brand
TEAL_600 = RGBColor.from_string("0A7C7C")
TEAL_100 = RGBColor.from_string("C4E0E0")
TEAL_50  = RGBColor.from_string("E5F2F2")
SUCCESS_700 = RGBColor.from_string("059669")
SUCCESS_50  = RGBColor.from_string("D1FAE5")
WARNING_700 = RGBColor.from_string("B45309")
WARNING_50  = RGBColor.from_string("FEF3C7")
DANGER_600  = RGBColor.from_string("DC2626")
DANGER_50   = RGBColor.from_string("FEE2E2")
INFO_700    = RGBColor.from_string("1D4ED8")
INFO_50     = RGBColor.from_string("DBEAFE")
N_50  = RGBColor.from_string("F7F8F9")
N_100 = RGBColor.from_string("ECEEF0")
N_200 = RGBColor.from_string("D9DFE3")
N_300 = RGBColor.from_string("B3BEC9")
N_500 = RGBColor.from_string("667685")
TEXT  = RGBColor.from_string("243342")
WHITE = RGBColor.from_string("FFFFFF")

FONT = "Segoe UI"

SLIDE_W = 13.333
SLIDE_H = 7.5
MARGIN = 0.6
CONTENT_W = SLIDE_W - 2 * MARGIN


# ── low-level helpers ────────────────────────────────────────────────────────
def _shape(slide, kind, x, y, w, h, *, fill=None, line=None, line_w=0.75,
           dash=None, radius=None):
    sp = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    sp.shadow.inherit = False
    if radius is not None:
        try:
            sp.adjustments[0] = radius
        except Exception:
            pass
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid()
        sp.fill.fore_color.rgb = fill
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line
        sp.line.width = Pt(line_w)
        if dash is not None:
            sp.line.dash_style = dash
    return sp


def rect(slide, x, y, w, h, **kw):
    return _shape(slide, MSO_SHAPE.RECTANGLE, x, y, w, h, **kw)


def rrect(slide, x, y, w, h, radius=0.08, **kw):
    return _shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, radius=radius, **kw)


def arrow(slide, x, y, w, h, color=N_300):
    return _shape(slide, MSO_SHAPE.RIGHT_ARROW, x, y, w, h, fill=color)


def text(slide, x, y, w, h, paras, *, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
         wrap=True):
    """paras: list of dicts. Either {'runs': [run, ...]} or a single-run dict.
    Run keys: text, size, color, bold, italic, font. Para keys: align, sb, sa, ls."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, p in enumerate(paras):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = p.get("align", align)
        if p.get("sb"):
            para.space_before = Pt(p["sb"])
        if p.get("sa"):
            para.space_after = Pt(p["sa"])
        para.line_spacing = p.get("ls", 1.08)
        for r in p.get("runs", [p]):
            run = para.add_run()
            run.text = r.get("text", "")
            f = run.font
            f.name = r.get("font", FONT)
            f.size = Pt(r.get("size", 14))
            f.bold = r.get("bold", False)
            f.italic = r.get("italic", False)
            f.color.rgb = r.get("color", TEXT)
    return tb


def P(t, size=14, color=TEXT, bold=False, italic=False, **kw):
    d = {"text": t, "size": size, "color": color, "bold": bold, "italic": italic}
    d.update(kw)
    return d


# ── deck builder ─────────────────────────────────────────────────────────────
class Deck:
    def __init__(self, filename: str, deck_label: str, role_title: str):
        self.prs = Presentation()
        self.prs.slide_width = Inches(SLIDE_W)
        self.prs.slide_height = Inches(SLIDE_H)
        self.filename = filename
        self.deck_label = deck_label
        self.role_title = role_title
        self.n = 0

    # -- scaffolding ---------------------------------------------------------
    def _blank(self):
        slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        self.n += 1
        return slide

    def _footer(self, slide, dark_bg=False):
        c = TEAL_100 if dark_bg else N_300
        text(slide, MARGIN, SLIDE_H - 0.42, 8.0, 0.3,
             [P(f"UniOps EPMS  ·  {self.deck_label}", 9, c)])
        text(slide, SLIDE_W - 1.6, SLIDE_H - 0.42, 1.0, 0.3,
             [P(str(self.n), 9, c)], align=PP_ALIGN.RIGHT)

    def _content(self, title, kicker=None):
        """Standard content slide chrome; returns slide and content top (in)."""
        slide = self._blank()
        rect(slide, 0, 0, SLIDE_W, 0.14, fill=TEAL_700)
        y = 0.45
        if kicker:
            text(slide, MARGIN, y, CONTENT_W, 0.3,
                 [P(kicker.upper(), 11, TEAL_600, bold=True)])
            y += 0.34
        text(slide, MARGIN, y, CONTENT_W, 0.6, [P(title, 26, TEXT, bold=True)])
        rect(slide, MARGIN, y + 0.62, 0.55, 0.05, fill=TEAL_700)
        self._footer(slide)
        return slide, y + 0.95

    # -- slide types ---------------------------------------------------------
    def cover(self, subtitle: str, audience: str):
        slide = self._blank()
        rect(slide, 0, 0, SLIDE_W, SLIDE_H, fill=TEAL_700)
        rect(slide, 0, SLIDE_H - 0.9, SLIDE_W, 0.9, fill=TEAL_800)
        text(slide, MARGIN + 0.4, 1.35, 9.5, 0.4,
             [P("UNIOPS EPMS  ·  ROLE TRAINING", 13, TEAL_100, bold=True)])
        text(slide, MARGIN + 0.4, 2.0, 11.5, 1.9,
             [P(self.role_title, 52, WHITE, bold=True)])
        text(slide, MARGIN + 0.4, 3.6, 11.0, 0.6, [P(subtitle, 19, TEAL_50)])
        rrect(slide, MARGIN + 0.4, 4.55, 5.9, 0.52, radius=0.5, fill=TEAL_600)
        text(slide, MARGIN + 0.66, 4.66, 5.6, 0.35,
             [P(f"Audience: {audience}", 12.5, WHITE, bold=True)])
        text(slide, MARGIN + 0.4, SLIDE_H - 0.68, 11.0, 0.35,
             [P("Enterprise Procurement Management  ·  Internal training material",
                11, TEAL_100)])

    def agenda(self, items: list[str]):
        slide, top = self._content("What This Training Covers")
        y = top + 0.15
        for i, it in enumerate(items, 1):
            rrect(slide, MARGIN, y, 0.42, 0.42, radius=0.5, fill=TEAL_50)
            text(slide, MARGIN, y + 0.055, 0.42, 0.32,
                 [P(str(i), 14, TEAL_700, bold=True)], align=PP_ALIGN.CENTER)
            text(slide, MARGIN + 0.62, y + 0.05, CONTENT_W - 0.8, 0.4,
                 [P(it, 16, TEXT)])
            y += 0.62

    def section(self, num: str, title: str, sub: str = ""):
        slide = self._blank()
        rect(slide, 0, 0, SLIDE_W, SLIDE_H, fill=N_50)
        rect(slide, 0, 0, 0.22, SLIDE_H, fill=TEAL_700)
        text(slide, MARGIN + 0.4, 2.15, 3.0, 1.5, [P(num, 88, TEAL_100, bold=True)])
        text(slide, MARGIN + 0.5, 3.55, 11.0, 0.9, [P(title, 36, TEXT, bold=True)])
        if sub:
            text(slide, MARGIN + 0.5, 4.35, 10.5, 0.6, [P(sub, 15, N_500)])
        self._footer(slide)

    def bullets(self, title, items, kicker=None, note=None, size=15, gap=0.14):
        """items: str or (head, desc)."""
        slide, top = self._content(title, kicker)
        y = top + 0.1
        for it in items:
            head, desc = (it, None) if isinstance(it, str) else it
            rrect(slide, MARGIN + 0.02, y + 0.09, 0.11, 0.11, radius=0.5, fill=TEAL_600)
            text(slide, MARGIN + 0.32, y, CONTENT_W - 0.5, 0.4,
                 [P(head, size, TEXT, bold=desc is not None)])
            y += 0.36 if len(head) < 95 else 0.62
            if desc:
                text(slide, MARGIN + 0.32, y, CONTENT_W - 0.7, 0.4,
                     [P(desc, size - 2.5, N_500)])
                y += (0.32 if len(desc) < 105 else 0.55) + gap
            else:
                y += gap
        if note:
            self._note(slide, note, y + 0.12)
        return slide

    def _note(self, slide, note_text, y, w=CONTENT_W):
        h = 0.62 if len(note_text) < 120 else 0.88
        y = min(y, SLIDE_H - 0.55 - h)
        rrect(slide, MARGIN, y, w, h, radius=0.12, fill=WARNING_50)
        text(slide, MARGIN + 0.22, y + 0.1, w - 0.45, h - 0.2,
             [{"runs": [P("Note:  ", 12, WARNING_700, bold=True),
                        P(note_text, 12, TEXT)]}])

    def two_col(self, title, left_title, left_items, right_title, right_items,
                kicker=None, note=None):
        slide, top = self._content(title, kicker)
        col_w = (CONTENT_W - 0.4) / 2
        for cx, (ct, items) in ((MARGIN, (left_title, left_items)),
                                (MARGIN + col_w + 0.4, (right_title, right_items))):
            rrect(slide, cx, top, col_w, 0.46, radius=0.14, fill=TEAL_50)
            text(slide, cx + 0.2, top + 0.09, col_w - 0.4, 0.32,
                 [P(ct, 14, TEAL_800, bold=True)])
            y = top + 0.66
            for it in items:
                head, desc = (it, None) if isinstance(it, str) else it
                rrect(slide, cx + 0.05, y + 0.085, 0.1, 0.1, radius=0.5, fill=TEAL_600)
                text(slide, cx + 0.28, y, col_w - 0.4, 0.4,
                     [P(head, 13, TEXT, bold=desc is not None)])
                y += 0.32 if len(head) < 55 else 0.56
                if desc:
                    text(slide, cx + 0.28, y, col_w - 0.42, 0.5,
                         [P(desc, 11, N_500)])
                    y += (0.28 if len(desc) < 62 else 0.5) + 0.08
                else:
                    y += 0.08
        if note:
            self._note(slide, note, 6.35)
        return slide

    def flow(self, title, steps, highlight, kicker=None, legend=None, below=None):
        """steps: list of labels; highlight: set of indices owned by this role."""
        slide, top = self._content(title, kicker)
        n = len(steps)
        gap = 0.34
        bw = (CONTENT_W - gap * (n - 1)) / n
        bh = 0.95
        y = top + 0.25
        for i, label in enumerate(steps):
            x = MARGIN + i * (bw + gap)
            hot = i in highlight
            rrect(slide, x, y, bw, bh, radius=0.12,
                  fill=TEAL_700 if hot else WHITE,
                  line=None if hot else N_200, line_w=1.1)
            text(slide, x + 0.05, y, bw - 0.1, bh,
                 [P(label, 11.5, WHITE if hot else TEXT, bold=hot)],
                 align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
            if i < n - 1:
                arrow(slide, x + bw + 0.05, y + bh / 2 - 0.07, gap - 0.1, 0.14)
        yy = y + bh + 0.28
        if legend:
            text(slide, MARGIN, yy, CONTENT_W, 0.3,
                 [{"runs": [P("■ ", 12, TEAL_700), P(legend, 12.5, N_500)]}])
            yy += 0.42
        if below:
            for head, desc in below:
                rrect(slide, MARGIN + 0.02, yy + 0.09, 0.11, 0.11, radius=0.5,
                      fill=TEAL_600)
                text(slide, MARGIN + 0.32, yy, CONTENT_W - 0.5, 0.4,
                     [P(head, 14, TEXT, bold=True)])
                yy += 0.34
                text(slide, MARGIN + 0.32, yy, CONTENT_W - 0.7, 0.4,
                     [P(desc, 12, N_500)])
                yy += 0.42
        return slide

    def pills(self, title, seq, kicker=None, rows_below=None, note=None):
        """seq: list of (label, bg, fg). Drawn as a pill chain with arrows.
        rows_below: list of (label, bg, fg, desc) legend lines."""
        slide, top = self._content(title, kicker)
        n = len(seq)
        gap = 0.3
        pw = min(1.75, (CONTENT_W - gap * (n - 1)) / n)
        total = pw * n + gap * (n - 1)
        x = MARGIN + (CONTENT_W - total) / 2
        y = top + 0.3
        for i, (label, bg, fg) in enumerate(seq):
            rrect(slide, x, y, pw, 0.52, radius=0.5, fill=bg)
            text(slide, x, y + 0.1, pw, 0.34, [P(label, 12, fg, bold=True)],
                 align=PP_ALIGN.CENTER)
            if i < n - 1:
                arrow(slide, x + pw + 0.04, y + 0.19, gap - 0.08, 0.14)
            x += pw + gap
        yy = y + 0.95
        if rows_below:
            for label, bg, fg, desc in rows_below:
                rrect(slide, MARGIN, yy, 1.75, 0.42, radius=0.5, fill=bg)
                text(slide, MARGIN, yy + 0.07, 1.75, 0.3,
                     [P(label, 11.5, fg, bold=True)], align=PP_ALIGN.CENTER)
                text(slide, MARGIN + 2.0, yy + 0.07, CONTENT_W - 2.1, 0.35,
                     [P(desc, 12.5, TEXT)])
                yy += 0.58
        if note:
            self._note(slide, note, yy + 0.1)
        return slide

    def steps_shot(self, title, steps, shot_label, kicker=None, caption=None):
        """Numbered steps on the left, screenshot placeholder on the right."""
        slide, top = self._content(title, kicker)
        y = top + 0.05
        for i, (head, desc) in enumerate(steps, 1):
            rrect(slide, MARGIN, y, 0.36, 0.36, radius=0.5, fill=TEAL_700)
            text(slide, MARGIN, y + 0.045, 0.36, 0.28,
                 [P(str(i), 12.5, WHITE, bold=True)], align=PP_ALIGN.CENTER)
            text(slide, MARGIN + 0.54, y - 0.01, 6.1, 0.35,
                 [P(head, 13.5, TEXT, bold=True)])
            y += 0.32
            text(slide, MARGIN + 0.54, y, 6.0, 0.55, [P(desc, 11.5, N_500)])
            y += (0.3 if len(desc) < 68 else 0.52) + 0.13
        sx, sy, sw, sh = 7.35, top + 0.05, 5.35, 4.9
        rrect(slide, sx, sy, sw, sh, radius=0.04, fill=N_50, line=N_300,
              line_w=1.2, dash=MSO_LINE.DASH)
        text(slide, sx, sy + sh / 2 - 0.55, sw, 0.5,
             [P("SCREENSHOT PLACEHOLDER", 13, N_300, bold=True)],
             align=PP_ALIGN.CENTER)
        text(slide, sx + 0.35, sy + sh / 2 - 0.05, sw - 0.7, 0.9,
             [P(shot_label, 12, N_500)], align=PP_ALIGN.CENTER)
        if caption:
            text(slide, sx, sy + sh + 0.08, sw, 0.5, [P(caption, 10.5, N_500,
                 italic=True)], align=PP_ALIGN.CENTER)
        return slide

    def table(self, title, headers, rows, kicker=None, col_widths=None, note=None,
              body_size=11.5, row_h=0.42):
        slide, top = self._content(title, kicker)
        n_rows, n_cols = len(rows) + 1, len(headers)
        h = min(0.5 + row_h * len(rows), 5.4)
        gf = slide.shapes.add_table(n_rows, n_cols, Inches(MARGIN), Inches(top),
                                    Inches(CONTENT_W), Inches(h))
        tbl = gf.table
        if col_widths:
            for i, w in enumerate(col_widths):
                tbl.columns[i].width = Inches(w)
        for j, htxt in enumerate(headers):
            cell = tbl.cell(0, j)
            cell.fill.solid()
            cell.fill.fore_color.rgb = TEAL_700
            cell.margin_left = cell.margin_right = Inches(0.12)
            tf = cell.text_frame
            tf.word_wrap = True
            pr = tf.paragraphs[0]
            run = pr.add_run()
            run.text = htxt
            run.font.name = FONT
            run.font.size = Pt(12)
            run.font.bold = True
            run.font.color.rgb = WHITE
        for i, row in enumerate(rows, 1):
            for j, val in enumerate(row):
                cell = tbl.cell(i, j)
                cell.fill.solid()
                cell.fill.fore_color.rgb = WHITE if i % 2 else N_50
                cell.margin_left = cell.margin_right = Inches(0.12)
                tf = cell.text_frame
                tf.word_wrap = True
                pr = tf.paragraphs[0]
                run = pr.add_run()
                run.text = str(val)
                run.font.name = FONT
                run.font.size = Pt(body_size)
                run.font.bold = j == 0
                run.font.color.rgb = TEXT if j == 0 else N_500
        if note:
            self._note(slide, note, top + h + 0.25)
        return slide

    def faq(self, qas, title="FAQ & Common Mistakes"):
        slide, top = self._content(title)
        y = top
        for q, a in qas:
            text(slide, MARGIN, y, CONTENT_W, 0.35,
                 [{"runs": [P("Q  ", 13.5, TEAL_700, bold=True),
                            P(q, 13.5, TEXT, bold=True)]}])
            y += 0.34
            text(slide, MARGIN + 0.36, y, CONTENT_W - 0.6, 0.5,
                 [P(a, 12, N_500)])
            y += (0.3 if len(a) < 110 else 0.52) + 0.17
        return slide

    def closing(self, lines):
        slide = self._blank()
        rect(slide, 0, 0, SLIDE_W, SLIDE_H, fill=TEAL_700)
        text(slide, MARGIN + 0.4, 1.7, 11.5, 1.0,
             [P("Questions & Support", 40, WHITE, bold=True)])
        y = 3.1
        for head, desc in lines:
            rrect(slide, MARGIN + 0.42, y + 0.1, 0.13, 0.13, radius=0.5, fill=TEAL_100)
            text(slide, MARGIN + 0.75, y, 10.8, 0.4, [P(head, 16, WHITE, bold=True)])
            y += 0.4
            text(slide, MARGIN + 0.75, y, 10.8, 0.4, [P(desc, 13, TEAL_50)])
            y += 0.55
        text(slide, MARGIN + 0.4, SLIDE_H - 0.75, 11.0, 0.4,
             [P("Thank you — welcome to UniOps EPMS.", 14, TEAL_100)])

    def save(self):
        path = os.path.join(OUT_DIR, self.filename)
        core = self.prs.core_properties
        core.title = self.deck_label
        core.author = "UniOps EPMS"
        self.prs.save(path)
        print(f"  {self.filename}: {self.n} slides")


# ── shared content ───────────────────────────────────────────────────────────
P2P_STEPS = ["PR\nRequisition", "PR\nApproval", "PO\nPurchase Order", "Order\nPlaced",
             "GR\nGoods Receipt", "Invoice", "PA\nPayment Application", "Payment"]

GETTING_AROUND = [
    ("Sign in through the UniOps Portal",
     "Open the UniOps Portal, sign in with your company account, then open the EPMS module card."),
    ("Sidebar — MY WORKSPACE",
     "Dashboard (your role-specific overview), Task Inbox (everything waiting for you), Reports."),
    ("Sidebar — PROCUREMENT",
     "Purchase Requisitions, Purchase Orders, Goods Receipt, Invoices, Payment Applications — you only see what your role can access."),
    ("Task Inbox is your to-do list",
     "Every action the system expects from you appears here with the document number, amount and days waiting. Work from it daily."),
    ("Notifications",
     "You also receive email / Teams notifications with a direct link to the task. Manage preferences in My Profile."),
]


def slide_getting_around(d: Deck, extra=None):
    items = list(GETTING_AROUND)
    if extra:
        items.extend(extra)
    d.bullets("Getting Around EPMS", items, kicker="Basics")


HELP_LINES = [
    ("Your first stop: the Task Inbox", "If the system expects something from you, it is listed there."),
    ("Document questions", "Check the Approval Timeline and Document Chain on the document detail page first."),
    ("Access or permission issues", "Contact your system administrator to review your role and permissions."),
    ("Process questions", "Reach out to the Procurement / Finance team leads or your department manager."),
]


# ═════════════════════════════════════════════════════════════════════════════
# Deck 1 — Requester
# ═════════════════════════════════════════════════════════════════════════════
def build_requester():
    d = Deck("EPMS_Training_Requester.pptx", "Requester Training", "Requester")

    d.cover("Raising purchase requisitions and confirming what you receive",
            "All staff who request goods or services")

    d.agenda([
        "Your role in the Procure-to-Pay flow",
        "Getting around EPMS — Portal, sidebar, Task Inbox",
        "Choosing the right procurement type",
        "Creating and submitting a Purchase Requisition (PR)",
        "Prepayment Required — flagging vendors that need payment up front",
        "Budget check, approval chain and document statuses",
        "Tracking your documents",
        "Confirming receipt — Goods Receipt, collection and service confirmation",
        "FAQ and common mistakes",
    ])

    d.flow("Your Role in the Procure-to-Pay Flow", P2P_STEPS, {0, 4},
           kicker="Overview",
           legend="Highlighted steps are yours. Procurement and Finance handle the rest.",
           below=[
               ("You start the chain", "Every purchase begins with your PR. A clear, complete PR is the fastest route to an approved order."),
               ("You confirm the end of delivery", "When goods arrive or a service is completed, you confirm receipt so invoices can be matched and paid."),
           ])

    slide_getting_around(d)

    d.bullets("Task Inbox — What Lands on Your Desk", [
        ("revise_pr — a PR was returned to you",
         "An approver sent your PR back with comments. Open it, fix what was asked, and resubmit."),
        ("collect_goods — parcel ready for pick-up",
         "Small deliveries are received centrally; you get a task to collect and confirm."),
        ("confirm_service_gr — confirm a completed service",
         "For service POs you confirm the work was done. This has an SLA — reminders escalate if you sit on it."),
        ("acknowledge_gr — acknowledge a receipt",
         "Confirm that goods received on your behalf are correct and undamaged."),
    ], kicker="Basics",
        note="Tasks also arrive as email / Teams notifications with a direct link — but the Inbox is the single complete list.")

    d.section("01", "Creating a Purchase Requisition",
              "From choosing the type to pressing Submit")

    d.table("Step Zero — Pick the Right Procurement Type",
            ["Type", "What it covers", "Notes"],
            [
                # Keep in step with epms-api/app/knowledge/pr_types.yaml — the
                # assistant answers from that file, and a deck that says
                # something different teaches people the wrong thing. See PRD
                # §2.1. "Extra scrutiny in approval" used to sit on the Type 5
                # row and was never true: approval routing reads department
                # configuration and cross-department payments, and has never
                # read procurement type.
                ["Type 1 — Raw Mat. / Packaging", "Production raw materials and packaging", "Not available in EPMS — handled outside this system"],
                ["Type 2 — Misc / Consumables", "Office supplies, small equipment, consumables", "Most common for everyday requests"],
                ["Type 3 — Spare Parts", "Maintenance and repair parts", "Pick the part from the Parts Catalog per line"],
                ["Type 4 — Service", "Contracted services, maintenance work, consulting", "Needs an expected completion date; you confirm the work yourself"],
                ["Type 5 — Fixed Asset", "Capitalizable equipment and assets", "Needs a Fixed Asset ID before you can submit — ask Finance"],
                ["Type 6 — Project-Related", "Purchases charged to a project", "Needs a Project No. and a completion date; you confirm it yourself"],
            ],
            kicker="Create a PR",
            col_widths=[3.4, 5.0, 3.73],
            note="The type drives which fields appear, how budget is checked and how receipt is confirmed — choose carefully.",
            body_size=11, row_h=0.5)

    d.steps_shot("Create a PR — Header", [
        ("Go to Purchase Requisitions → New PR", "Or use the New PR shortcut on your Dashboard."),
        ("Select the Procurement Type", "The card selector at the top. Type 1 is disabled by design."),
        ("Fill in the header", "Department, currency, needed-by date and the business purpose. For Type 6, pick the Project."),
        ("Vendor needs payment up front? Tick 'Prepayment Required'", "Deposit or 100% before delivery — this one checkbox routes the whole payment flow (next slide)."),
        ("Write a purpose people can approve", "One or two sentences: what, why, and for whom. Vague purposes get returned."),
    ], "PR Create page — procurement type selector and header fields\n(/pr/new)",
        kicker="Create a PR",
        caption="Replace with a live screenshot of the New PR page.")

    d.steps_shot("Create a PR — Line Items", [
        ("Add one line per item", "Description, quantity, unit (pcs, kg, hour, …) and unit price. Line totals calculate automatically."),
        ("Types 2 & 3: link master data", "Select the material / part where applicable; add the supplier item # (e.g. an Amazon ASIN) if you have it."),
        ("Prices are estimates", "Use your best quote. Procurement fixes the final price on the PO."),
        ("Use notes for specifics", "Color, model, delivery constraints — anything the buyer must know."),
    ], "PR Create page — line items grid with totals\n(/pr/new)",
        kicker="Create a PR")

    d.flow("Prepayment Required — Small Checkbox, Big Consequences",
           ["PR\nPrepayment\nRequired ☑", "PO\nPrepaid flag\npre-filled", "Prepayment PA\nVendor paid\nbefore delivery", "Delivery\nYou confirm\nreceipt", "Settlement PA\nPrepayment\ncleared"],
           {0, 3},
           kicker="Create a PR",
           legend="Your checkbox starts this chain; your receipt confirmation unlocks the final settlement.",
           below=[
               ("When to tick it", "The vendor's quote asks for a deposit or payment before shipping — attach the quote showing those terms."),
               ("Why it matters", "Without the flag AP has no upfront-payment route — the vendor won't ship while payment is reworked."),
               ("Settlement is on a clock", "Overdue settlements escalate to management and block PO closure — confirm your receipt promptly."),
           ])

    d.bullets("Budget Check Before You Submit", [
        ("The Budget Balance widget shows live availability",
         "As you fill in the PR, it shows the budget line affected and the remaining balance after your request."),
        ("Insufficient budget = friction in approval",
         "Approvers see the same numbers. If the balance is not enough, expect a return or rejection."),
        ("Fix budget problems before submitting",
         "Talk to your Department Manager or Finance BP about the budget plan — do not just submit and hope."),
    ], kicker="Create a PR")

    d.bullets("Attachments & Submit", [
        ("Attach your evidence", "Vendor quotes, spec sheets, screenshots of the item — approvers approve faster when they can see what they are approving."),
        ("Save Draft vs Submit", "Draft keeps the PR private and editable. Submit sends it into the approval chain and locks editing."),
        ("After submit", "You can follow progress on the PR detail page. You will be notified at each approval step outcome."),
    ], kicker="Create a PR")

    d.section("02", "Approval & Tracking", "Who approves, what the statuses mean, where to look")

    d.flow("The PR Approval Chain",
           ["You\nSubmit", "Supervisor", "Department\nManager", "Director", "GM / OPM", "Approved"],
           {0},
           kicker="Approval",
           legend="You submit; the chain routes automatically based on your department.",
           below=[
               ("Steps can be skipped", "Supervisor and Director steps depend on your department's configuration — the timeline always shows your actual chain."),
               ("Approvers act from anywhere", "Approvals happen on the web, by email or in Teams — you do not need to chase people down."),
           ])

    d.pills("Document Statuses — What They Mean",
            [("Draft", N_100, TEXT), ("Submitted", INFO_50, INFO_700),
             ("In Review", WARNING_50, WARNING_700), ("Approved", SUCCESS_50, SUCCESS_700)],
            kicker="Approval",
            rows_below=[
                ("Returned", WARNING_50, WARNING_700,
                 "Sent back to you for changes — fix and resubmit. This is normal, not a failure."),
                ("Rejected", DANGER_50, DANGER_600,
                 "Definitively declined. Start a new PR if the need still exists."),
                ("Cancelled", N_100, TEXT,
                 "Withdrawn — the PR is closed and will not be processed."),
            ])

    d.steps_shot("Tracking Your PRs", [
        ("PR List — your documents at a glance", "Filter by status, search by number. You see your own PRs (and your department's, if configured)."),
        ("Approval Timeline", "The detail page shows every step: who, when, through which channel, and any comments."),
        ("Document Chain", "See everything downstream of your PR — the PO, receipts, invoices and payments it produced."),
        ("Returned? Check the comment first", "The approver's comment tells you exactly what to change before you resubmit."),
    ], "PR Detail page — Approval Timeline and Document Chain\n(/pr/{id})",
        kicker="Tracking")

    d.section("03", "Confirming Receipt", "Goods Receipt, collection and service confirmation")

    d.bullets("Goods Receipt (GR) — Why It Matters", [
        ("No confirmed receipt, no payment",
         "Finance matches invoices against received quantities. Unconfirmed deliveries block vendor payment."),
        ("Goods deliveries", "Warehouse or the requester records what physically arrived against the PO. Partial deliveries are normal — receive what came."),
        ("Report problems at receipt", "Damaged or wrong items should be flagged during receipt so Procurement can follow up with the vendor."),
    ], kicker="Receiving")

    d.bullets("Collection & Service Confirmation", [
        ("Collection confirmation",
         "Small parcels are received centrally. You get a collect_goods task — pick up the parcel and confirm collection in the task."),
        ("Service confirmation",
         "For service POs (Type 4) you confirm the service was actually completed — this is the service equivalent of a goods receipt."),
        ("Service confirmations have an SLA",
         "If you do not confirm in time you will be reminded, then escalated. Confirm promptly — or dispute promptly if the work is not done."),
    ], kicker="Receiving",
        note="Only confirm what is genuinely delivered or completed. Your confirmation authorizes payment.")

    d.faq([
        ("My PR was returned. Do I have to start over?",
         "No. Open the revise_pr task, edit the PR, address the approver's comment and resubmit — it re-enters the approval chain."),
        ("Can I edit a PR after submitting?",
         "Not while it is in review. Either wait for the outcome or ask the current approver to return it to you."),
        ("The budget widget shows insufficient funds. Can I still submit?",
         "Technically yes, but expect a rejection. Resolve the budget question with your Department Manager / Finance BP first."),
        ("The vendor wants a 50% deposit before shipping.",
         "Tick Prepayment Required on the PR and attach the quote. AP pays the deposit up front and settles after delivery."),
        ("Where do I see whether my order has been placed or delivered?",
         "Open your PR and use the Document Chain — it links to the PO and receipts with their live statuses."),
        ("I received goods but never created a GR. Is that a problem?",
         "Yes — the vendor's invoice cannot be matched and paid. Create or confirm the receipt as soon as goods arrive."),
    ])

    d.closing(HELP_LINES)
    d.save()


# ═════════════════════════════════════════════════════════════════════════════
# Deck 2 — Procurement Officer
# ═════════════════════════════════════════════════════════════════════════════
def build_procurement_officer():
    d = Deck("EPMS_Training_Procurement_Officer.pptx",
             "Procurement Officer Training", "Procurement Officer")

    d.cover("Turning approved requisitions into orders — and getting them delivered",
            "Procurement Officers and Procurement Managers")

    d.agenda([
        "Your role in the Procure-to-Pay flow",
        "Getting around EPMS — Dashboard and Task Inbox",
        "Creating a Purchase Order from an approved PR",
        "PO pricing, tax codes and prepayment flag",
        "PO approval chain and statuses",
        "Placing the order — email and online",
        "Delivery follow-up and Goods Receipt",
        "Master data you own — Vendors, Parts Catalog, Projects",
        "FAQ and common mistakes",
    ])

    d.flow("Your Role in the Procure-to-Pay Flow", P2P_STEPS, {2, 3},
           kicker="Overview",
           legend="Highlighted steps are yours — plus delivery follow-up until the PO is fully received.",
           below=[
               ("You convert demand into commitment", "An approved PR becomes a legally meaningful PO with real prices, taxes and a chosen vendor."),
               ("You own the vendor relationship", "Placing orders, chasing deliveries and keeping vendor master data clean all sit with you."),
           ])

    slide_getting_around(d, extra=[
        ("Sidebar — MASTER DATA",
         "Vendors, Projects and the Parts Catalog are maintained by the procurement team."),
    ])

    d.bullets("Task Inbox — What Lands on Your Desk", [
        ("create_po — an approved PR needs a buyer",
         "A PR finished its approval chain. Review it and turn it into a Purchase Order."),
        ("place_order — an approved PO is ready to send",
         "The PO passed approval. Send it to the vendor by email or place the order online."),
        ("revise_po — a PO was returned to you",
         "Procurement Manager or Finance Manager sent it back. Fix and resubmit."),
        ("gr_damage_report — a receipt flagged problems",
         "Goods arrived damaged or wrong. Follow up with the vendor for replacement or credit."),
    ], kicker="Basics")

    d.section("01", "From PR to PO", "Creating and submitting the Purchase Order")

    d.steps_shot("Create a PO from an Approved PR", [
        ("Start from the task or PO → New PO", "The create_po task links straight to the source PR."),
        ("Select the vendor", "Pick from the vendor master. If the vendor does not exist yet, create it first (or import from ERP)."),
        ("Confirm lines and real prices", "PR prices are estimates. Enter the quoted prices, adjust quantities and units where needed."),
        ("Check currency and delivery details", "Currency, delivery address and terms print on the PO document the vendor receives."),
    ], "PO Create page — source PR, vendor selector and line grid\n(/po/new)",
        kicker="PR to PO",
        caption="Replace with a live screenshot of the New PO page.")

    d.bullets("Pricing, Tax and Prepayment", [
        ("Tax codes come from master data",
         "Select the tax code per the vendor / goods situation. The rate is snapshotted onto the PO — later master-data changes do not alter existing POs."),
        ("Never hand-type tax rates",
         "If a rate is missing or wrong, have Finance fix it in Tax Settings — do not work around it."),
        ("Prepaid PO flag",
         "Pre-filled from the PR's Prepayment Required flag — verify it against the vendor's actual terms. Payment Applications on a prepaid PO default to the prepayment flow."),
        ("Totals are what gets approved",
         "Net, tax and gross totals are computed live — check them against the vendor quote before submitting."),
    ], kicker="PR to PO")

    d.flow("The PO Approval Chain",
           ["You\nSubmit", "Procurement\nManager", "Finance\nManager", "Approved", "Place\nOrder"],
           {0, 4},
           kicker="Approval",
           legend="Two approval levels: commercial review, then financial review.",
           below=[
               ("Withdraw while you still can", "A draft or submitted PO can be withdrawn if you spot a mistake — fix it and resubmit."),
               ("The PO PDF is generated on approval", "Once approved, the official PO document is created automatically and attached to the PO."),
           ])

    d.pills("PO Lifecycle — Statuses",
            [("Draft", N_100, TEXT), ("Submitted", INFO_50, INFO_700),
             ("In Review", WARNING_50, WARNING_700), ("Approved", SUCCESS_50, SUCCESS_700),
             ("Issued", TEAL_50, TEAL_800)],
            kicker="Approval",
            rows_below=[
                ("Partially Received", INFO_50, INFO_700, "Some lines/quantities have been received — keep chasing the rest."),
                ("Fully Received", SUCCESS_50, SUCCESS_700, "Everything delivered. Invoices can be matched against the full order."),
                ("Returned / Rejected", DANGER_50, DANGER_600, "Returned = fix and resubmit. Rejected = the order will not proceed."),
            ])

    d.section("02", "Placing the Order", "Getting the approved PO to the vendor")

    d.steps_shot("Place Order — Two Ways", [
        ("Open the approved PO", "The 'Ready to Place Order' panel appears for approved POs — or work from the place_order task."),
        ("Option A — Place Order via Email", "Send the PO PDF to the vendor directly from EPMS using the company email template. Adjust recipients and message as needed."),
        ("Option B — Place Order Online", "Ordered through a vendor portal or webshop instead? Record it here so the PO still moves to Issued."),
        ("PO becomes Issued", "Either way the PO is now with the vendor and delivery tracking begins."),
    ], "PO Detail — Place Order modal with Email / Online options\n(/po/{id})",
        kicker="Ordering",
        caption="Replace with a live screenshot of the Place Order modal.")

    d.bullets("After Issuing — Delivery Follow-Up", [
        ("Track receipt progress on the PO",
         "Issued → Partially Received → Fully Received, driven by Goods Receipts recorded against your PO."),
        ("Receipts are recorded by warehouse / requesters",
         "You do not usually create GRs yourself, but you monitor them — the PO detail shows received vs ordered per line."),
        ("Damage reports come to you",
         "A gr_damage_report task means the receipt flagged problems. Contact the vendor for replacement, credit note or return."),
        ("Service POs",
         "The requester confirms service completion (with an SLA). Chase the vendor if work stalls; chase the requester if confirmation stalls."),
    ], kicker="Delivery")

    d.section("03", "Master Data You Own", "Vendors, Parts Catalog and Projects")

    d.two_col("Vendors",
              "Keeping the vendor master clean",
              [
                  ("One vendor, one record", "Search before you create — duplicate vendors fragment spend history and confuse invoice matching."),
                  ("Complete the profile", "Legal name, contacts, payment details and default currency — the PO email and payment runs rely on them."),
                  ("ERP import", "Vendors can be imported from the ERP to avoid retyping and keep identifiers aligned."),
              ],
              "When to touch it",
              [
                  ("New vendor needed for a PO", "Create or import the vendor before creating the PO."),
                  ("Vendor details changed", "Update the record — do not carry corrections in PO notes."),
                  ("Deactivate, don't delete", "Retired vendors keep their history but disappear from pickers."),
              ],
              kicker="Master data")

    d.two_col("Parts Catalog & Projects",
              "Parts Catalog (Type 3 backbone)",
              [
                  ("Requesters pick parts from here", "A well-maintained catalog means accurate spare-parts PRs."),
                  ("Add parts with real identifiers", "Part number, description, unit — vague entries produce vague orders."),
              ],
              "Projects (Type 6 backbone)",
              [
                  ("Project list drives cost tracking", "Type 6 PRs must reference a project; keep the list current."),
                  ("Close finished projects", "So new spend cannot be booked against them."),
              ],
              kicker="Master data")

    d.faq([
        ("The PR's prices don't match the vendor's quote. Which wins?",
         "Yours. The PR is an estimate; you enter the real quoted prices on the PO. Large deviations deserve a comment for the approvers."),
        ("Can I change a PO after it was approved?",
         "No — approved POs are commitments. Material changes require a new or revised order; talk to your Procurement Manager."),
        ("I submitted a PO with a mistake. What now?",
         "Withdraw it (possible while draft/submitted), fix it, resubmit. Once in review, ask the approver to return it instead."),
        ("Where is the PO PDF?",
         "Generated automatically when the PO is approved — check the attachments panel on the PO detail page."),
        ("The vendor needs payment before delivery.",
         "Mark the PO as prepaid. AP then raises a prepayment PA, and a settlement follows after delivery."),
        ("The tax rate looks wrong on my PO.",
         "Tax codes are managed by Finance in Tax Settings. Get the master data fixed — the PO snapshots whatever is current at creation."),
    ])

    d.closing(HELP_LINES)
    d.save()


# ═════════════════════════════════════════════════════════════════════════════
# Deck 3 — AP Clerk
# ═════════════════════════════════════════════════════════════════════════════
def build_ap_clerk():
    d = Deck("EPMS_Training_AP_Clerk.pptx", "AP Clerk Training", "AP Clerk")

    d.cover("From vendor invoice to executed payment — matching, applications and settlement",
            "Accounts Payable clerks")

    d.agenda([
        "Your role in the Procure-to-Pay flow",
        "Getting around EPMS — Dashboard and Task Inbox",
        "Invoice intake — upload and OCR verification",
        "Matching invoices to POs — single and multi-PO allocation",
        "Exceptions, duplicates and invoice statuses",
        "Payment Applications — the four types",
        "Prepayment and settlement in depth",
        "PA approval chain and payment execution",
        "FAQ and common mistakes",
    ])

    d.flow("Your Role in the Procure-to-Pay Flow", P2P_STEPS, {5, 6},
           kicker="Overview",
           legend="Highlighted steps are yours — through to monitoring payment execution.",
           below=[
               ("You are the control point before money leaves", "Nothing gets paid without an invoice you registered and matched, and a PA you raised."),
               ("Accuracy over speed", "A wrong match or a duplicate invoice costs far more time downstream than careful checking costs now."),
           ])

    slide_getting_around(d, extra=[
        ("Finance visibility",
         "Your role also opens the Portal's Finance views (Budget Dashboard, Budget Plans, Finance) for cross-checking."),
    ])

    d.bullets("Task Inbox — What Lands on Your Desk", [
        ("link_invoice — an invoice needs matching",
         "An uploaded invoice is still unmatched. Match it to its PO(s) so it can proceed."),
        ("create_pa — a payment application is due",
         "Deliveries or terms indicate it is time to request payment for a PO."),
        ("settle_prepayment — a prepayment awaits settlement",
         "Goods/services for a prepaid PO were delivered — raise the Settlement PA to clear the prepayment."),
        ("process_pa / revise_pa — PA follow-ups",
         "Returned PAs come back to you with comments; fix and resubmit."),
    ], kicker="Basics")

    d.section("01", "Invoice Intake", "Upload, OCR pre-fill and verification")

    d.steps_shot("Uploading Vendor Invoices", [
        ("Go to Invoices → Upload", "Drag & drop the PDF (or browse). You can upload several files in one go."),
        ("OCR pre-fills the fields", "Vendor, invoice number, dates and amounts are extracted automatically from the PDF."),
        ("Verify every pre-filled field", "OCR is an assistant, not an authority. Compare against the PDF preview side by side."),
        ("Check tax lines", "Confirm the tax amounts and codes match what is printed on the invoice."),
    ], "Invoice List — upload dropzone with OCR-prefilled form and PDF preview\n(/invoices)",
        kicker="Intake",
        caption="Replace with a live screenshot of the invoice upload flow.")

    d.bullets("Before You Save — the Intake Checklist", [
        ("Right vendor?", "OCR can pick a similarly-named vendor. The vendor drives matching and payment — get it right."),
        ("Invoice number exactly as printed", "Vendor + invoice number identifies the invoice. Typos create phantom duplicates later."),
        ("Amounts reconcile", "Net + tax = gross, matching the PDF to the cent."),
        ("Already in the system?", "Search the invoice list first — the same invoice uploaded twice is the classic AP error."),
    ], kicker="Intake")

    d.section("02", "Matching", "Connecting invoices to purchase orders")

    d.steps_shot("Match an Invoice to Its PO", [
        ("Open the unmatched invoice", "From the link_invoice task or the Invoices list filtered to Unmatched."),
        ("Pick the PO", "Search by PO number or vendor. The system shows the PO's ordered, received and already-invoiced amounts."),
        ("Confirm the match", "Amounts within tolerance → the invoice moves to Matched and becomes available for a PA."),
        ("Deviation? Don't force it", "If amounts do not reconcile, treat it as an exception — investigate before matching."),
    ], "Invoice Detail — PO match panel\n(/invoices/{id})",
        kicker="Matching")

    d.bullets("One Invoice, Several POs — Line Allocation", [
        ("When a vendor bills several orders on one invoice",
         "Use the allocation panel instead of a single match — split the invoice across POs at line level."),
        ("Allocate every dollar",
         "The allocations must add up to the invoice total; the panel shows the running remainder."),
        ("Each PO sees its share",
         "Downstream, every PO carries exactly its allocated portion for payment and reporting."),
    ], kicker="Matching",
        note="Allocation is the correct tool for consolidated invoices — never register the same invoice twice to cover two POs.")

    d.pills("Invoice Lifecycle — Statuses",
            [("Unmatched", WARNING_50, WARNING_700), ("Matched", SUCCESS_50, SUCCESS_700),
             ("Approved", INFO_50, INFO_700), ("Paid", TEAL_50, TEAL_800)],
            kicker="Matching",
            rows_below=[
                ("Exception", DANGER_50, DANGER_600,
                 "Something does not reconcile — price, quantity or a missing PO. Resolve it explicitly; exceptions are worked, not ignored."),
            ],
            note="Only unmatched / exception invoices can be deleted. Once matched, an invoice is part of the audit trail.")

    d.section("03", "Payment Applications", "Requesting, approving and executing payment")

    d.table("The Four PA Types",
            ["Type", "When to use it", "Key fields"],
            [
                ["Regular", "Standard payment after goods/services are received and the invoice is matched", "PO, paid lines & quantities, linked invoice"],
                ["Prepayment", "Vendor requires payment before delivery (prepaid PO)", "Prepayment % and expected settlement date"],
                ["Settlement", "Delivery complete on a prepaid PO — settle against the prepayment", "Original prepayment PA, prepayment applied, net payable"],
                ["Balance", "Remaining balance related to an earlier prepayment arrangement", "Original prepayment PA reference"],
            ],
            kicker="Payment Applications",
            col_widths=[1.9, 6.0, 4.23],
            body_size=11, row_h=0.55)

    d.steps_shot("Create a Regular PA", [
        ("Start from the task or PA → New PA", "Select the PO — its lines, received quantities and prior PAs load automatically."),
        ("Choose what you are paying", "Pick lines and quantities for this application (partial payments are fine)."),
        ("Link the matched invoice", "The PA should be backed by the vendor invoice you matched earlier."),
        ("Verify totals and submit", "Net, tax and gross must agree with the invoice. Then submit into approval."),
    ], "PA Create page — PO selector, line picker, totals\n(/pa/new)",
        kicker="Payment Applications")

    d.bullets("Prepayment & Settlement — the Full Cycle", [
        ("1. Prepayment PA",
         "On a prepaid PO, raise a Prepayment PA: set the percentage (e.g. 50%) and the expected settlement date. Vendor gets paid before delivery."),
        ("2. Delivery happens",
         "Goods arrive / service completes and is confirmed like any other PO."),
        ("3. Settlement PA — use the Settle button",
         "Open the prepayment PA and click Settle. The PO, the original prepayment and the applied amount are pre-filled and locked."),
        ("4. Net payable = invoice total − prepayment applied",
         "Only the remainder is paid out. If the prepayment covers everything, net is zero — a pure reconciliation with no cash movement."),
        ("Watch the settlement clock",
         "Overdue settlements trigger reminders and escalate. The settle_prepayment task is your cue."),
    ], kicker="Payment Applications", size=13.5, gap=0.1)

    d.flow("The PA Approval Chain",
           ["You\nSubmit", "Department\nManager", "Director", "GM / OPM", "Finance\nBP", "Finance\nManager", "Approved"],
           {0},
           kicker="Approval & payment",
           legend="Routing follows the department of the original PR's creator; steps may be skipped by configuration.",
           below=[
               ("After approval — payment execution", "Finance executes the payment through the unified payment run; the PA (and its invoice) move to Paid."),
               ("Returned PAs come back to you", "A revise_pa task with the approver's comment — adjust and resubmit."),
           ])

    d.faq([
        ("OCR filled in the wrong amount and I saved it. Now what?",
         "Edit the invoice while it is still unmatched — verification against the PDF is exactly why the preview sits next to the form."),
        ("I uploaded the same invoice twice.",
         "Delete the duplicate while it is unmatched. Matched invoices are part of the audit trail and cannot simply be deleted."),
        ("The invoice is more than what the PO has left.",
         "Do not force the match. Check for a second PO (use allocation), a price change, or a vendor error — resolve as an exception."),
        ("The invoice has no PO at all.",
         "Non-PO spend does not go through EPMS invoice matching — route it per the finance process (e.g. OA direct payment)."),
        ("A settlement nets to zero. Do I still create it?",
         "Yes. A zero-net settlement reconciles the prepayment against the delivery — no cash moves, and the prepayment is closed properly."),
        ("When is the vendor actually paid?",
         "After the PA clears its full approval chain, Finance executes payment in the payment run; statuses then flip to Paid automatically."),
    ])

    d.closing(HELP_LINES)
    d.save()


if __name__ == "__main__":
    print("Generating EPMS training decks ->", OUT_DIR)
    build_requester()
    build_procurement_officer()
    build_ap_clerk()
    print("Done.")
