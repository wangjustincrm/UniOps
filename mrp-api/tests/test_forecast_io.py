"""Sales-forecast Excel template/export (Phase 1A Task 2).

`POST /versions/{id}/import` — this file's original subject — was retired
in the Continuous Sales Forecast redesign (Task 8) along with the rest of
the per-version write/draft API (see tests/test_forecast.py's module
docstring); `forecast_io.py`'s import-parsing helpers
(`parse_import_workbook`, `fetch_valid_material_codes`, `ImportFileError`)
have no remaining HTTP caller. Only template/export coverage remains here,
built against an ORM fixture version rather than via the retired
POST/PUT/confirm endpoints (mirrors tests/test_forecast.py's `_make_version`
helper).
"""
import io
import uuid
from decimal import Decimal

import pytest
from openpyxl import load_workbook

from app.models.forecast import ForecastLine, ForecastVersion


async def _make_version(db_session, start_month="2026-09", horizon_months=18) -> ForecastVersion:
    version = ForecastVersion(
        version_no=f"FCV-{start_month}-{uuid.uuid4().hex[:8].upper()}",
        status="draft",
        horizon_start_month=start_month,
        horizon_months=horizon_months,
    )
    db_session.add(version)
    await db_session.commit()
    await db_session.refresh(version)
    return version


@pytest.mark.anyio
async def test_template_download_has_material_name_and_18_month_headers(client, db_session, admin_token):
    headers = {"Authorization": f"Bearer {admin_token}"}
    v = await _make_version(db_session)
    months = (await client.get(f"/api/v1/forecast/versions/{v.id}/grid", headers=headers)).json()["months"]

    r = await client.get(f"/api/v1/forecast/versions/{v.id}/template", headers=headers)
    assert r.status_code == 200
    assert "spreadsheetml" in r.headers["content-type"]
    assert v.version_no in r.headers["content-disposition"]

    wb = load_workbook(io.BytesIO(r.content))
    ws = wb.active
    header = [c.value for c in next(ws.iter_rows(max_row=1))]
    assert header == ["Material Code", "Name"] + months
    assert ws.max_row == 1  # empty template — no data rows


@pytest.mark.anyio
async def test_export_reflects_current_grid(client, db_session, admin_token):
    headers = {"Authorization": f"Bearer {admin_token}"}
    v = await _make_version(db_session)
    months = (await client.get(f"/api/v1/forecast/versions/{v.id}/grid", headers=headers)).json()["months"]
    db_session.add(ForecastLine(version_id=v.id, material_code="S0093", month=months[0], qty=Decimal("42")))
    await db_session.commit()

    r = await client.get(f"/api/v1/forecast/versions/{v.id}/export", headers=headers)
    assert r.status_code == 200
    wb = load_workbook(io.BytesIO(r.content))
    ws = wb.active
    header = [c.value for c in next(ws.iter_rows(max_row=1))]
    assert header == ["Material Code", "Name"] + months

    data_rows = list(ws.iter_rows(min_row=2, values_only=True))
    assert len(data_rows) == 1
    assert data_rows[0][0] == "S0093"
    assert float(data_rows[0][2]) == 42
