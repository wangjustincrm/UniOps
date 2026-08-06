"""Sales forecast Excel template / import / export helpers (Phase 1A Task 2).

Consumes Task 1's `ForecastVersion`/`ForecastLine` shape via the caller
(app/api/v1/forecast.py) — this module only deals with xlsx bytes and
row-level validation, it does not touch the DB session itself.

`fetch_valid_material_codes` is deliberately a *synchronous* function that
pages through mdm-api's `GET /mdm/v1/materials` **once per import call**
(never per row) forwarding the caller's bearer token. It is sync so tests
can monkeypatch this exact module attribute with a plain callable (see
tests/test_forecast_io.py, mirroring the task brief's
`monkeypatch.setattr(forecast_io, "fetch_valid_material_codes", lambda *a,
**k: {...})`) without needing an async mock; the import endpoint bridges it
onto the event loop via `anyio.to_thread.run_sync`. It now derives from
`app.services.mdm_client.fetch_materials` — the same httpx.Client/paging
loop, generalized to also carry each material's name so read endpoints
(forecast grid/export, consignment stock list, net-requirement) can resolve
`material_code -> name` without a second, independent way of calling
mdm-api.

Row numbering in `error_rows`/ok-cell reporting matches the **physical Excel
row number** the user would see if they opened the file — the header is row
1, so the first data row is row 2. Users read this report to locate the bad
row in their spreadsheet, so it must line up with what Excel shows, not a
0-based or header-excluded count. A header-level problem (mismatched month
column) is reported as row 0 (there is no single data row it belongs to).

**Hardening (I7, final-phase review):** the caller (app/api/v1/forecast.py's
`import_forecast`) enforces file size/extension/content-type before this
module ever sees the bytes, but the actual xlsx parse still has two failure
modes this module owns:

- `load_workbook(..., read_only=True)` — read-only mode streams rows lazily
  instead of materializing the whole parsed workbook in memory, so an xlsx
  crafted to decompress to something far larger than it looks on disk (a
  "zip bomb") can't OOM the container just from being opened, even under
  `dry_run=true` where nothing is ever written.
- A non-xlsx or corrupt file (e.g. a `.csv` renamed to `.xlsx`, or a
  genuinely truncated upload) makes `load_workbook` raise — openpyxl's own
  exception, `zipfile.BadZipFile`, or something else entirely depending on
  how the bytes are malformed. `parse_import_workbook` catches all of that
  and re-raises as `ImportFileError`, a `ValueError` subclass with a
  human-readable message, so the endpoint can turn it into a 400 instead of
  an opaque 500.

`MAX_IMPORT_DATA_ROWS` caps the number of data rows this module will walk
per import, independent of the byte-size cap the endpoint enforces (a small
file can still claim an enormous number of rows) — exceeding it raises
`ImportFileError` rather than silently truncating, so the user knows to
split the file instead of unknowingly importing a partial grid.
"""
from __future__ import annotations

import io
from decimal import Decimal, InvalidOperation

from openpyxl import Workbook, load_workbook

from app.services.mdm_client import fetch_materials

TEMPLATE_HEADERS = ["Material Code", "Name"]

# Hard cap on data rows a single import will walk (see module docstring's
# "Hardening" section) — independent of the endpoint's byte-size cap.
MAX_IMPORT_DATA_ROWS = 20_000


class ImportFileError(ValueError):
    """The uploaded file itself is unreadable as an xlsx workbook, or blows
    past a hard structural limit (too many rows) — as opposed to a
    per-row/per-cell content problem, which is reported via `error_rows`
    instead of raising. The endpoint (app/api/v1/forecast.py) catches this
    and returns 400 with `str(exc)` as the detail."""


def fetch_valid_material_codes(token: str) -> set[str]:
    """Pull the full set of material codes known to mdm-api, once.

    Thin wrapper over `app.services.mdm_client.fetch_materials` (see this
    module's docstring) — the codes are just that map's keys. Forwards the
    caller's bearer token so this respects mdm-api's own authz (materials
    reads there are open to any authenticated role). Raises on failure
    (never falls back to an empty set) — see `fetch_materials`'s docstring
    for why import_forecast depends on that.
    """
    return set(fetch_materials(token).keys())


def parse_qty(raw) -> Decimal:
    """Tolerant, non-negative quantity parser.

    Strips thousands separators (`,`) and `$` before parsing, so `"20,000"`
    and `"$1,234.50"` both parse cleanly. Raises ValueError (with a
    human-readable message) for blank/non-numeric/negative input.
    """
    if raw is None:
        raise ValueError("quantity is required")
    text = str(raw).strip()
    if text == "":
        raise ValueError("quantity is required")
    cleaned = text.replace(",", "").replace("$", "").strip()
    try:
        value = Decimal(cleaned)
    except InvalidOperation as exc:
        raise ValueError(f"'{raw}' is not a valid number") from exc
    if value < 0:
        raise ValueError(f"quantity must be non-negative, got '{raw}'")
    return value


def build_template_workbook(months: list[str]) -> bytes:
    """Empty template: header row only — Material Code | Name | <18 months>."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Forecast"
    ws.append(TEMPLATE_HEADERS + list(months))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_export_workbook(months: list[str], rows: list[dict]) -> bytes:
    """Same shape as the template, filled with the current grid.

    `rows`: [{"material_code": str, "name": str | None, "cells": {month: Decimal}}]
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Forecast"
    ws.append(TEMPLATE_HEADERS + list(months))
    for row in rows:
        cells = row.get("cells") or {}
        values = [row["material_code"], row.get("name") or ""]
        for month in months:
            qty = cells.get(month)
            # openpyxl accepts Decimal natively (it's a real numeric cell
            # type, not coerced to text) — going through float() here was
            # the one place float ever entered this chain (M11, final-phase
            # review) and could shift the last digit on an export -> import
            # round trip, which test_export_then_reimport_round_trips_clean
            # depends on being exact.
            values.append(qty if qty is not None else "")
        ws.append(values)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def parse_import_workbook(
    file_bytes: bytes, months: list[str], valid_codes: set[str],
) -> tuple[list[dict], list[dict]]:
    """Validate an uploaded xlsx against `months`/`valid_codes`.

    Returns (ok_cells, error_rows):
      - ok_cells: [{"material_code", "month", "qty": Decimal}] ready for the
        same upsert semantics as PUT .../cells.
      - error_rows: [{"row", "column", "reason"}]. `row` is the physical
        Excel row number (header = row 1, first data row = row 2). One row's
        problem never stops the rest — every other row is still evaluated.

    Raises `ImportFileError` (see module docstring) if the bytes aren't a
    readable xlsx workbook at all, or if it has more than
    `MAX_IMPORT_DATA_ROWS` data rows — both are structural problems with the
    file itself, not a per-row content problem, so they don't belong in
    `error_rows`.
    """
    try:
        # read_only=True: stream rows instead of materializing the whole
        # parsed workbook — see module docstring's "Hardening" section for
        # why (a zip-bomb-style xlsx must not be able to OOM the container
        # just from being opened, even under dry_run=true).
        wb = load_workbook(io.BytesIO(file_bytes), data_only=True, read_only=True)
    except Exception as exc:
        # openpyxl raises different exception types for different kinds of
        # bad input (InvalidFileException for a recognizably-wrong format,
        # zipfile.BadZipFile for a non-zip/corrupt file, KeyError for a
        # zip that's missing an expected part, ...) — catch broadly and
        # normalize to one readable message rather than letting whichever
        # one comes up surface as an opaque 500.
        raise ImportFileError(
            f"could not read the uploaded file as an Excel (.xlsx) workbook: {exc}"
        ) from exc
    ws = wb.active

    rows_iter = ws.iter_rows(values_only=True)
    try:
        header = list(next(rows_iter) or ())
    except StopIteration:
        return [], [{"row": 0, "column": "", "reason": "workbook has no header row"}]

    header_months = header[2:2 + len(months)]
    header_months = ["" if h is None else str(h).strip() for h in header_months]

    error_rows: list[dict] = []
    valid_month_positions: list[int] = []
    for i, expected in enumerate(months):
        actual = header_months[i] if i < len(header_months) else ""
        if actual != expected:
            error_rows.append({
                "row": 0,
                "column": actual or f"(column {i + 3})",
                "reason": f"expected month header '{expected}', got '{actual or '(missing)'}'",
            })
        else:
            valid_month_positions.append(i)

    ok_cells: list[dict] = []
    # start=2: the header consumed by next() above is Excel row 1, so the
    # first row this loop sees is Excel row 2 — report the real row number.
    for row_num, raw_row in enumerate(rows_iter, start=2):
        # row_num - 1 == the number of data rows seen so far (row 2 is data
        # row 1). Checked before any per-row work so a file that blows past
        # the cap fails fast rather than after fully scanning e.g. a
        # 500k-row sheet.
        if row_num - 1 > MAX_IMPORT_DATA_ROWS:
            raise ImportFileError(
                f"the uploaded file has more than {MAX_IMPORT_DATA_ROWS} data rows "
                "— split it into smaller files and import them separately"
            )
        if raw_row is None:
            continue
        row = list(raw_row)
        material_code = str(row[0]).strip() if row and row[0] is not None else ""
        if material_code == "":
            continue  # blank row — nothing to import

        if material_code not in valid_codes:
            error_rows.append({
                "row": row_num,
                "column": "Material Code",
                "reason": f"unknown material code '{material_code}'",
            })
            continue

        for i in valid_month_positions:
            col_index = 2 + i
            cell_value = row[col_index] if col_index < len(row) else None
            if cell_value is None or str(cell_value).strip() == "":
                continue  # blank cell — no data for this material/month
            try:
                qty = parse_qty(cell_value)
            except ValueError as exc:
                error_rows.append({"row": row_num, "column": months[i], "reason": str(exc)})
                continue
            ok_cells.append({"material_code": material_code, "month": months[i], "qty": qty})

    return ok_cells, error_rows
