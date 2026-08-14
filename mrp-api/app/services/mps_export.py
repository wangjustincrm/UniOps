"""MPS production-plan matrix -> xlsx (Production Plan Matrix design, Task 2;
week columns under a month grouping header, Task 8 of the 2026-08-12 weekly
rework).

Mirrors `app/services/forecast_io.py`'s export shape (openpyxl `Workbook`,
saved to an in-memory `io.BytesIO`, returned as raw bytes for the caller —
`app/api/v1/mps.py`'s `GET /runs/{id}/export` — to wrap in a `Response`).

Unlike the forecast grid (one row per material, one column per month, one
number per cell), the production-plan matrix needs FOUR numbers per
material/week — Demand, Available (opening stock), Planned, Gap — so each
product is rendered as four rows instead of one. Columns are WEEKS, not
`demand_month` (when the forecast wants it) and not `plan_week_month` either
(the shape this module had before Task 8) — a planner reading a *production*
plan needs to see the actual week the factory floor builds in, not just the
month it falls in. Row 1 groups those week columns under their owning month
(merged cells); row 2 carries each week's `week_label`; data starts at row 3.

**Columns come from the run's horizon, not from which weeks the lines
happened to land on.** `weeks_of_month` is walked for every month from
`run.horizon_start_month` for `run.horizon_months`, UNION every
`plan_week_month` actually present on `lines` (a lead-shifted pre-build can
legitimately land a line in a month before the nominal horizon start — see
`app/api/v1/mps.py::_planning_weeks`'s docstring — and dropping that column
would silently lose the line's data), spanned contiguously from the earliest
to the latest month so no month in between is skipped either. A week with no
line landed on it still gets a column of zeros rather than not existing —
the time axis must not silently skip.

Week labels are rendered under `run.week_calendar_mode` — the run's OWN
stored mode, never the live planning parameter (see `app/api/v1/mps.py`'s
module docstring, "A run snapshots the calendar it was generated under") —
so a released plan's export looks the same today as it will next year even
if somebody changes the factory's week-boundary convention in between.

`lines` are expected to already carry each line's demand context
(`demand_forecast`/`opening_stock`) — the caller builds these via
`app/api/v1/mps.py`'s `_build_demand_context`/`_line_response` (Task 1 of
this design), exactly the same computation `GET /runs/{id}` uses. This
module only aggregates and renders; it never re-derives net-requirement math
itself.

Decimal stays Decimal all the way from `MrpMpsLine.qty` / the demand context
through the per-(material, week) aggregation — the ONLY place a value is
coerced to `float` is the openpyxl cell boundary in `_scaled`, mirroring
`forecast_io.build_export_workbook`'s own "float only at the very edge"
discipline (see that module's docstring, M11 finding).

## Capacity-gap lines get their own row, excluded from Planned

A `capacity_gap` line is an un-placed shortfall pinned to the week it was
supposed to be made in — `app/services/mps_engine.py`'s docstring is explicit
that it "never touches the ledger" and is not a booked production slot.
Task 7's review flagged that the export nonetheless summed a gap line's
`qty` straight into "Planned", i.e. reported unmet demand as if it had been
built. That's wrong on its face — a planner exporting the plan to hand to
the factory floor must see production, not a number that quietly includes
what could not be produced.

Fixed here by giving Gap its own metric row instead of folding it into
Planned or dropping it. A silently-excluded gap is *safer* than counting it
as output but still hides the shortfall from a planner who only skims
Planned vs. Demand; an explicit Gap row keeps Planned honest (real production
only) while still surfacing exactly how much demand went unplaced and in
which week, so nothing needs cross-referencing the API response to see it.
The row is emitted for every product even when it is all zero, so the sheet
shape is stable across exports/products instead of coming and going with
whether a run happens to have any shortfalls.

## Demand and Available are MONTHLY, spanned across the month's weeks

Only **Planned** and **Gap** are per-week. `demand_forecast`/`opening_stock`
are snapshotted PER DEMAND MONTH and copied onto every line of that month
(mrp07), so they are month-grain figures with no weekly meaning at all —
design §5.1 says so explicitly ("Demand 与 Available 仍按月……显示在该月的
第一周列并跨列居中；只有 Planned 落到具体周").

They are therefore aggregated per `(material_code, plan_week_month)`,
written into the month's FIRST week column, and merged across exactly the
span row 1 already merges the month header over — one span, computed once
in `month_columns` and used by both, so the header and the values cannot
drift apart. This is what `ProductionMatrix.tsx`'s `MonthMetricCell` does on
the other side of the wire (`colSpan={monthSpans.get(month)}`, reading the
month-grain `aggregateLines(lines, monthCellKey(code, plan_week_month))`
map), and the two must agree: the same run rendered two ways must not tell
a planner two different things.

Writing the month figure into every week column instead — which this module
did between Task 8 and the final review — made the sheet read Demand
120/120/120/120 against Planned 30/30/30/30, i.e. a 90 t weekly shortfall
that does not exist, and made the Demand row sum to 4x the real demand for
anyone who selected it in Excel.

Within a month the per-demand-month snapshot is still counted ONCE per
distinct `(material_code, plan_week_month, demand_month)` (the `counted`
set) and never once per line: one demand month routinely spans several
lines, and summing them was the N x inflation Task 7 first fixed (measured
360 against a real 120). A demand month whose lines straddle TWO plan
months contributes once to EACH of those months — never zero times (it must
not vanish from either group) and never twice within one.
"""
from __future__ import annotations

import io
from datetime import date
from decimal import Decimal, ROUND_HALF_UP
from typing import Iterable, Protocol

from openpyxl import Workbook
from openpyxl.styles import Alignment

from app.services.week_calendar import week_label, weeks_of_month

_KG_PER_TONNE = Decimal("1000")
_THREE_DP = Decimal("0.001")

# Row order within each product's four-row block.
_METRIC_ROW_LABELS = ("Demand", "Available", "Planned", "Gap")

# The two month-grain metrics: written once per MONTH in that month's first
# week column and merged across it (design 5.1), not once per week. The
# other two are per-week. See the module docstring's "Demand and Available
# are MONTHLY" section.
_MONTH_GRAIN_METRICS = frozenset({"Demand", "Available"})

_SPANNED = Alignment(horizontal="center", vertical="center")

# Data columns start at column 3 (1-indexed): col 1 = Product, col 2 = Metric.
_FIRST_DATA_COLUMN = 3


class _LineLike(Protocol):
    material_code: str
    demand_month: str
    plan_week_start: date
    plan_week_month: str
    qty: Decimal
    demand_forecast: Decimal
    opening_stock: Decimal
    capacity_gap: bool
    # Minimum lot size (mrp11): how much of this month was already covered
    # by an earlier batch's surplus. It is part of what is AVAILABLE to the
    # month, so the sheet must add it the same way the on-screen matrix
    # does -- the two are mirrored implementations and a divergence here
    # shows up as an export that contradicts the screen.
    carry_in_qty: Decimal


def _scaled(value: Decimal, unit: str) -> float:
    """Decimal KG -> the cell's numeric value. `unit == "t"` divides by 1000
    and rounds to 3dp (design brief); `unit == "kg"` passes the raw KG value
    through unchanged. This is the one and only float() coercion in this
    module — everything upstream (aggregation) stays exact Decimal."""
    if unit == "t":
        value = (value / _KG_PER_TONNE).quantize(_THREE_DP, rounding=ROUND_HALF_UP)
    return float(value)


def _generate_months(start_month: str, count: int) -> list[str]:
    """Mirrors `app/api/v1/forecast.py::_generate_months` and
    `app/api/v1/net_requirement.py::_generate_months` exactly — those two are
    the real precedent for duplicating this three-line month walk rather
    than importing it (net_requirement's own docstring says "Mirrors
    app/api/v1/forecast.py::_generate_months exactly"). Duplicated here so a
    services module does not depend on `app/api/v1/*`.

    (An earlier version of this docstring cited
    `app/services/mps_engine.py::_generate_months`. That function has never
    existed in this repo — the engine walks weeks, not months.)"""
    year, month = (int(p) for p in start_month.split("-"))
    months = []
    for i in range(count):
        m0 = month - 1 + i
        y = year + m0 // 12
        mm = m0 % 12 + 1
        months.append(f"{y:04d}-{mm:02d}")
    return months


def _month_span(first: str, last: str) -> list[str]:
    """Every 'YYYY-MM' from `first` to `last` inclusive. Mirrors
    `app/api/v1/mps.py::_month_span`, duplicated for the same reason as
    `_generate_months` above."""
    start = int(first[:4]) * 12 + int(first[5:7]) - 1
    end = int(last[:4]) * 12 + int(last[5:7]) - 1
    return [f"{i // 12:04d}-{i % 12 + 1:02d}" for i in range(start, end + 1)]


def build_mps_matrix_workbook(
    run, lines: Iterable[_LineLike], unit: str, name_by_code: dict[str, str | None],
) -> bytes:
    """Group `lines` by material_code -> plan_week_start (the week columns)
    for Planned/Gap, and by material_code -> plan_week_month for
    Demand/Available.

    `planned` sums every non-gap line's qty landing on that exact week.
    `gap` sums every gap line's qty on that week (excluded from `planned` —
    see module docstring). `demand`/`available` are the frozen per-demand-
    month snapshot, month-grain, added once per distinct
    `(material_code, plan_week_month, demand_month)` — see the module
    docstring's "Demand and Available are MONTHLY" section for why they are
    not per week, and what the sheet looked like when they were.

    Sheet layout: row 1 = month grouping (merged across that month's week
    columns), row 2 = week labels (`week_label` under `run.week_calendar_mode`
    — the run's OWN stored mode), data from row 3: four rows per product
    (sorted by material_code) in Demand / Available / Planned / Gap order.
    Demand and Available occupy one merged, centred cell per month, over the
    SAME span as that month's row-1 header; Planned and Gap get one cell per
    week column.
    `name_by_code` resolves the Product cell to the material's display name,
    falling back to the bare code when the map has no entry (mdm-api
    degrade-to-{} contract — see `app.services.mdm_client.resolve_material_names`'s
    docstring).

    `run` supplies `run_no` (sheet title), `horizon_start_month`/
    `horizon_months` (the nominal column span) and `week_calendar_mode` (how
    those columns are labelled and how a line's own week is resolved to a
    column — see module docstring).
    """
    lines = list(lines)
    mode = run.week_calendar_mode
    # The run's OWN grid, never the current parameter: an exported plan must
    # have the same columns it had when it was released.
    start_dow = run.week_start_dow

    horizon_months = _generate_months(run.horizon_start_month, run.horizon_months)
    touched_months = {line.plan_week_month for line in lines}
    all_months = set(horizon_months) | touched_months
    span_months = _month_span(min(all_months), max(all_months)) if all_months else []

    # Ordered (month, week_start) pairs -- one per data column. A week
    # belongs to exactly one month under `weeks_of_month`'s own partition
    # (see week_calendar.py's docstring), so `week_start` is a safe unique
    # key across the whole grid.
    week_grid: list[tuple[str, date]] = [
        (month, w) for month in span_months
        for w in weeks_of_month(month, mode, start_dow=start_dow)
    ]

    # Each month's column span, 1-indexed and inclusive, derived from the
    # grid itself (the grid is month-contiguous by construction above).
    # ONE span, used by row 1's month header merge AND by every product's
    # Demand/Available merge -- computing it twice would let the header and
    # the values it labels drift apart, which is the class of bug this
    # module keeps hitting.
    month_columns: dict[str, tuple[int, int]] = {}
    for offset, (month, _w) in enumerate(week_grid):
        col = _FIRST_DATA_COLUMN + offset
        first, _last = month_columns.get(month, (col, col))
        month_columns[month] = (first, col)

    # Planned/Gap are keyed (material_code, plan_week_start) -- per WEEK.
    # Demand/Available are keyed (material_code, plan_week_month) -- per
    # MONTH, deduped on `counted` because the snapshot is copied onto every
    # line of a demand month (see module docstring).
    planned: dict[tuple[str, date], Decimal] = {}
    gap: dict[tuple[str, date], Decimal] = {}
    demand: dict[tuple[str, str], Decimal] = {}
    available: dict[tuple[str, str], Decimal] = {}
    counted: dict[tuple[str, str], set[str]] = {}
    materials: set[str] = set()

    for line in lines:
        materials.add(line.material_code)
        week_key = (line.material_code, line.plan_week_start)
        month_key = (line.material_code, line.plan_week_month)

        if line.capacity_gap:
            gap[week_key] = gap.get(week_key, Decimal("0")) + line.qty
        else:
            planned[week_key] = planned.get(week_key, Decimal("0")) + line.qty

        seen = counted.setdefault(month_key, set())
        if line.demand_month not in seen:
            seen.add(line.demand_month)
            demand[month_key] = demand.get(month_key, Decimal("0")) + line.demand_forecast
            available[month_key] = (available.get(month_key, Decimal("0"))
                                    + line.opening_stock + line.carry_in_qty)

    wb = Workbook()
    ws = wb.active
    ws.title = run.run_no

    ws.cell(row=1, column=1, value="Product")
    ws.cell(row=1, column=2, value="Metric")
    ws.merge_cells(start_row=1, start_column=1, end_row=2, end_column=1)
    ws.merge_cells(start_row=1, start_column=2, end_row=2, end_column=2)

    # Row 2 (week labels), then row 1 (month grouping merged across that
    # month's own week columns) from the shared `month_columns` span.
    for offset, (_month, w) in enumerate(week_grid):
        ws.cell(row=2, column=_FIRST_DATA_COLUMN + offset,
                value=week_label(w, mode, start_dow=start_dow))
    for month, (start_col, end_col) in month_columns.items():
        cell = ws.cell(row=1, column=start_col, value=month)
        cell.alignment = _SPANNED
        if end_col > start_col:
            ws.merge_cells(start_row=1, start_column=start_col, end_row=1, end_column=end_col)

    row = 3
    tables = {"Demand": demand, "Available": available, "Planned": planned, "Gap": gap}
    for material_code in sorted(materials):
        product_name = name_by_code.get(material_code) or material_code
        for label in _METRIC_ROW_LABELS:
            table = tables[label]
            ws.cell(row=row, column=1, value=product_name)
            ws.cell(row=row, column=2, value=label)
            if label in _MONTH_GRAIN_METRICS:
                # One merged, centred cell per month -- the same span row 1
                # merged the month header over. The rest of the span is left
                # EMPTY by the merge rather than repeating the figure:
                # repeating it made a levelled month read as a weekly
                # shortfall and made the row sum to N x the real demand.
                for month, (start_col, end_col) in month_columns.items():
                    value = table.get((material_code, month), Decimal("0"))
                    cell = ws.cell(row=row, column=start_col, value=_scaled(value, unit))
                    cell.alignment = _SPANNED
                    if end_col > start_col:
                        ws.merge_cells(start_row=row, start_column=start_col,
                                       end_row=row, end_column=end_col)
            else:
                for offset, (_month, w) in enumerate(week_grid):
                    value = table.get((material_code, w), Decimal("0"))
                    ws.cell(row=row, column=_FIRST_DATA_COLUMN + offset,
                            value=_scaled(value, unit))
            row += 1

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
