"""Warehouse receiving report — what belongs on the sheet and what must not."""
import uuid as _uuid
from datetime import datetime, timedelta, timezone

import pytest
from sqlalchemy import select as _select
from sqlalchemy.ext.asyncio import AsyncSession, async_sessionmaker

from tests.test_gr import _create_gr, _make_issued_po

REPORT_URL = "/api/v1/gr/receiving-report"
EXPORT_URL = "/api/v1/gr/receiving-report/export"


@pytest.fixture(autouse=True)
def _quiet_notification_email(monkeypatch):
    """Same reason as test_gr.py: an SMTP retry outliving the test deadlocks
    the next test's DROP ... CASCADE."""
    async def _noop(to, subject, html, **kwargs):
        return None

    monkeypatch.setattr("app.services.email.send_email", _noop)


def _factory(engine):
    return async_sessionmaker(engine, class_=AsyncSession, expire_on_commit=False)


async def _stamp(engine, *, po_id=None, gr_id=None, po_fields=None, gr_fields=None,
                 pr_fields=None):
    """Set columns the API does not accept from a client (placed_at, source,
    received_at, the PR's department) so a row can be posed at a known date."""
    from app.models.gr import GoodsReceipt
    from app.models.po import PurchaseOrder
    from app.models.pr import PurchaseRequest

    async with _factory(engine)() as db:
        if po_fields and po_id:
            po = (await db.execute(
                _select(PurchaseOrder).where(PurchaseOrder.id == _uuid.UUID(po_id))
            )).scalar_one()
            for key, value in po_fields.items():
                setattr(po, key, value)
        if gr_fields and gr_id:
            gr = (await db.execute(
                _select(GoodsReceipt).where(GoodsReceipt.id == _uuid.UUID(gr_id))
            )).scalar_one()
            for key, value in gr_fields.items():
                setattr(gr, key, value)
        if pr_fields and po_id:
            po = (await db.execute(
                _select(PurchaseOrder).where(PurchaseOrder.id == _uuid.UUID(po_id))
            )).scalar_one()
            pr = (await db.execute(
                _select(PurchaseRequest).where(PurchaseRequest.id == po.pr_id)
            )).scalar_one()
            for key, value in pr_fields.items():
                setattr(pr, key, value)
        await db.commit()


def _row_for(payload, gr_number):
    return next(r for r in payload["items"] if r["gr_number"] == gr_number)


# ── The row itself ────────────────────────────────────────────────────────────

@pytest.mark.asyncio
async def test_received_line_appears_with_every_column(admin_client, test_engine):
    po, _ = await _make_issued_po(admin_client, test_engine, "VND-RPT-01")
    gr = await _create_gr(admin_client, po["id"])
    await _stamp(
        test_engine, po_id=po["id"], gr_id=gr["id"],
        po_fields={"placed_at": datetime(2026, 8, 3, 14, 0, tzinfo=timezone.utc)},
        gr_fields={
            "received_at": datetime(2026, 8, 13, 14, 0, tzinfo=timezone.utc),
            "received_by": "Warehouse Migs",
            "collected_at": datetime(2026, 8, 14, 14, 0, tzinfo=timezone.utc),
            "collected_by": "Kirk",
        },
        pr_fields={"department_name": "Engineering"},
    )

    resp = await admin_client.get(REPORT_URL)
    assert resp.status_code == 200, resp.text
    row = _row_for(resp.json(), gr["number"])

    assert row["description"] == "Hydraulic Pump"
    assert row["supplier"] == "GR Vendor"
    assert row["po_number"] == po["number"]
    assert row["unit"] == "EA"
    assert row["quantity"] == "2.0000"
    assert row["department"] == "Engineering"
    assert row["requested_by"] == "GR Requester"
    assert row["date_ordered"] == "2026-08-03"
    assert row["arrival_date"] == "2026-08-13"
    assert row["left_warehouse_date"] == "2026-08-14"
    assert row["warehouse_receiver"] == "Warehouse Migs"
    assert row["person_accepting"] == "Kirk"
    assert row["lead_time_days"] == 10


@pytest.mark.asyncio
async def test_dates_are_plant_local_not_utc(admin_client, test_engine):
    """An evening receipt in Toronto is already the next day in UTC. Reporting
    the UTC day would move the row and shorten the lead time by one."""
    po, _ = await _make_issued_po(admin_client, test_engine, "VND-RPT-TZ")
    gr = await _create_gr(admin_client, po["id"])
    await _stamp(
        test_engine, po_id=po["id"], gr_id=gr["id"],
        # 2026-08-03 21:00 EDT and 2026-08-13 21:00 EDT — both 01:00 UTC the
        # following day.
        po_fields={"placed_at": datetime(2026, 8, 4, 1, 0, tzinfo=timezone.utc)},
        gr_fields={"received_at": datetime(2026, 8, 14, 1, 0, tzinfo=timezone.utc)},
    )

    row = _row_for((await admin_client.get(REPORT_URL)).json(), gr["number"])
    assert row["date_ordered"] == "2026-08-03"
    assert row["arrival_date"] == "2026-08-13"
    assert row["lead_time_days"] == 10


@pytest.mark.asyncio
async def test_date_ordered_falls_back_to_last_approval(admin_client, test_engine):
    """Half the orders on file predate the Place Order action. The last
    approval — when the order was cleared to go out — stands in for it, so the
    lead-time column is not blank for everything historical."""
    from app.models.approval import ApprovalEvent

    po, _ = await _make_issued_po(admin_client, test_engine, "VND-RPT-APPR")
    gr = await _create_gr(admin_client, po["id"])

    async with _factory(test_engine)() as db:
        from app.models.user import User
        actor = (await db.execute(_select(User).limit(1))).scalars().first()
        for step, day in ((0, 4), (1, 6)):
            db.add(ApprovalEvent(
                document_type="po", document_id=_uuid.UUID(po["id"]),
                document_number=po["number"], step_idx=step, action="approve",
                actor_id=actor.id, actor_role="gm",
                created_at=datetime(2026, 8, day, 14, 0, tzinfo=timezone.utc),
            ))
        await db.commit()

    await _stamp(
        test_engine, po_id=po["id"], gr_id=gr["id"],
        po_fields={"placed_at": None},
        gr_fields={"received_at": datetime(2026, 8, 16, 14, 0, tzinfo=timezone.utc)},
    )

    row = _row_for((await admin_client.get(REPORT_URL)).json(), gr["number"])
    assert row["date_ordered"] == "2026-08-06"    # the LAST approval, not the first
    assert row["lead_time_days"] == 10


@pytest.mark.asyncio
async def test_material_id_falls_back_to_the_ordered_line(admin_client, test_engine):
    """The receiver rarely re-types a material id; the ordered line has it."""
    from app.models.po import PoLineItem

    po, _ = await _make_issued_po(admin_client, test_engine, "VND-RPT-MAT")
    async with _factory(test_engine)() as db:
        line = (await db.execute(
            _select(PoLineItem).where(PoLineItem.po_id == _uuid.UUID(po["id"]))
        )).scalars().first()
        line.material_id = "M1637"
        await db.commit()
        po_line_id = str(line.id)

    gr = await _create_gr(admin_client, po["id"], line_items=[{
        "po_line_id": po_line_id,
        "description": "Hydraulic Pump", "qty_ordered": "2", "qty_received": "2",
        "unit": "EA", "unit_price": "350.00", "condition": "good",
    }])

    row = _row_for((await admin_client.get(REPORT_URL)).json(), gr["number"])
    assert row["material_id"] == "M1637"


# ── What the report must leave out ────────────────────────────────────────────

@pytest.mark.asyncio
async def test_service_confirmations_are_excluded(admin_client, test_engine):
    """A type-4 PO's GR confirms work done — there is nothing to receive."""
    po, _ = await _make_issued_po(admin_client, test_engine, "VND-RPT-SVC", po_type=4)
    gr = await _create_gr(admin_client, po["id"])
    assert gr["gr_type"] == "service"

    numbers = [r["gr_number"] for r in (await admin_client.get(REPORT_URL)).json()["items"]]
    assert gr["number"] not in numbers


@pytest.mark.asyncio
async def test_nc_mirrored_orders_are_excluded(admin_client, test_engine):
    """ERP-imported orders are received in NC and only mirrored here, so they
    are not part of the warehouse's own receiving log."""
    po, _ = await _make_issued_po(admin_client, test_engine, "VND-RPT-NC")
    gr = await _create_gr(admin_client, po["id"])

    before = [r["gr_number"] for r in (await admin_client.get(REPORT_URL)).json()["items"]]
    assert gr["number"] in before

    await _stamp(test_engine, po_id=po["id"], po_fields={"source": "nc"})
    after = [r["gr_number"] for r in (await admin_client.get(REPORT_URL)).json()["items"]]
    assert gr["number"] not in after


@pytest.mark.asyncio
async def test_cancelled_receipts_are_excluded(admin_client, test_engine):
    po, _ = await _make_issued_po(admin_client, test_engine, "VND-RPT-CXL")
    gr = await _create_gr(admin_client, po["id"])
    await _stamp(test_engine, gr_id=gr["id"], gr_fields={"status": "cancelled"})

    numbers = [r["gr_number"] for r in (await admin_client.get(REPORT_URL)).json()["items"]]
    assert gr["number"] not in numbers


# ── Filters ───────────────────────────────────────────────────────────────────

@pytest.mark.asyncio
async def test_date_window_is_inclusive_of_both_ends(admin_client, test_engine):
    po, _ = await _make_issued_po(admin_client, test_engine, "VND-RPT-WIN")
    gr = await _create_gr(admin_client, po["id"])
    # 2026-08-13 20:00 EDT — the last hours of the local day, and already the
    # 14th in UTC. An `date_to=2026-08-13` window must still contain it.
    await _stamp(test_engine, gr_id=gr["id"], gr_fields={
        "received_at": datetime(2026, 8, 14, 0, 0, tzinfo=timezone.utc),
    })

    inside = await admin_client.get(REPORT_URL, params={
        "date_from": "2026-08-13", "date_to": "2026-08-13"})
    assert gr["number"] in [r["gr_number"] for r in inside.json()["items"]]

    after = await admin_client.get(REPORT_URL, params={"date_from": "2026-08-14"})
    assert gr["number"] not in [r["gr_number"] for r in after.json()["items"]]

    before = await admin_client.get(REPORT_URL, params={"date_to": "2026-08-12"})
    assert gr["number"] not in [r["gr_number"] for r in before.json()["items"]]


@pytest.mark.asyncio
async def test_search_matches_the_line_description(admin_client, test_engine):
    po, _ = await _make_issued_po(admin_client, test_engine, "VND-RPT-SRCH")
    gr = await _create_gr(admin_client, po["id"], line_items=[{
        "description": "Balaclavas White Universal", "qty_ordered": "2",
        "qty_received": "2", "unit": "box", "unit_price": "10.00", "condition": "good",
    }])

    hit = await admin_client.get(REPORT_URL, params={"search": "Balaclava"})
    assert gr["number"] in [r["gr_number"] for r in hit.json()["items"]]

    miss = await admin_client.get(REPORT_URL, params={"search": "Nothing Like This"})
    assert gr["number"] not in [r["gr_number"] for r in miss.json()["items"]]


@pytest.mark.asyncio
async def test_department_filter(admin_client, test_engine):
    from app.models.department import Department

    po, _ = await _make_issued_po(admin_client, test_engine, "VND-RPT-DEPT")
    gr = await _create_gr(admin_client, po["id"])

    async with _factory(test_engine)() as db:
        dept = Department(code=f"D{_uuid.uuid4().hex[:6]}", name="Maintenance")
        db.add(dept)
        await db.commit()
        dept_id = str(dept.id)

    await _stamp(test_engine, po_id=po["id"],
                 pr_fields={"department_id": _uuid.UUID(dept_id), "department_name": "Maintenance"})

    hit = await admin_client.get(REPORT_URL, params={"department_id": dept_id})
    assert gr["number"] in [r["gr_number"] for r in hit.json()["items"]]

    other = await admin_client.get(REPORT_URL, params={"department_id": str(_uuid.uuid4())})
    assert gr["number"] not in [r["gr_number"] for r in other.json()["items"]]


# ── Export ────────────────────────────────────────────────────────────────────

@pytest.mark.asyncio
async def test_export_returns_a_workbook_with_the_same_rows(admin_client, test_engine):
    from io import BytesIO

    from openpyxl import load_workbook

    po, _ = await _make_issued_po(admin_client, test_engine, "VND-RPT-XLSX")
    gr = await _create_gr(admin_client, po["id"])
    await _stamp(
        test_engine, po_id=po["id"], gr_id=gr["id"],
        po_fields={"placed_at": datetime(2026, 8, 3, 14, 0, tzinfo=timezone.utc)},
        gr_fields={"received_at": datetime(2026, 8, 13, 14, 0, tzinfo=timezone.utc)},
    )

    resp = await admin_client.get(EXPORT_URL, params={
        "date_from": "2026-08-01", "date_to": "2026-08-31"})
    assert resp.status_code == 200, resp.text
    assert "spreadsheetml" in resp.headers["content-type"]
    assert "receiving-report" in resp.headers["content-disposition"]

    sheet = load_workbook(BytesIO(resp.content)).active
    headers = [c.value for c in sheet[1]]
    assert headers[:6] == [
        "Material ID", "Description", "Manufacturer/Supplier",
        "Purchase Order Number", "Unit of Measure", "Quantity",
    ]
    assert headers[13] == "Lead time"

    rows = [[c.value for c in r] for r in sheet.iter_rows(min_row=2)]
    mine = next(r for r in rows if r[14] == gr["number"])
    assert mine[1] == "Hydraulic Pump"
    assert mine[5] == 2.0                                   # a number, not a string
    assert mine[8].date().isoformat() == "2026-08-03"       # a date, not a string
    assert mine[9].date().isoformat() == "2026-08-13"
    assert mine[13] == 10


@pytest.mark.asyncio
async def test_report_is_empty_when_the_matrix_hides_grs(admin_client, test_engine, monkeypatch):
    """The report is a rendering of records the caller can already open — never
    a way around the Access Control Matrix."""
    po, _ = await _make_issued_po(admin_client, test_engine, "VND-RPT-AUTHZ")
    gr = await _create_gr(admin_client, po["id"])
    assert gr["number"] in [
        r["gr_number"] for r in (await admin_client.get(REPORT_URL)).json()["items"]
    ]

    real_build_scope = None
    from app.api.v1 import gr as gr_api

    async def _no_view(db, user):
        scope = await real_build_scope(db, user)
        scope["perms"]["view_gr"] = False
        return scope

    real_build_scope = gr_api.build_scope
    monkeypatch.setattr(gr_api, "build_scope", _no_view)

    payload = (await admin_client.get(REPORT_URL)).json()
    assert payload["items"] == []
    assert payload["total"] == 0


@pytest.mark.asyncio
async def test_rows_are_ordered_by_arrival_date(admin_client, test_engine):
    po_a, _ = await _make_issued_po(admin_client, test_engine, "VND-RPT-ORD-A")
    po_b, _ = await _make_issued_po(admin_client, test_engine, "VND-RPT-ORD-B")
    gr_late = await _create_gr(admin_client, po_a["id"])
    gr_early = await _create_gr(admin_client, po_b["id"])

    now = datetime.now(timezone.utc)
    await _stamp(test_engine, gr_id=gr_late["id"],
                 gr_fields={"received_at": now})
    await _stamp(test_engine, gr_id=gr_early["id"],
                 gr_fields={"received_at": now - timedelta(days=3)})

    numbers = [r["gr_number"] for r in (await admin_client.get(REPORT_URL)).json()["items"]]
    assert numbers.index(gr_early["number"]) < numbers.index(gr_late["number"])
