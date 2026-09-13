"""Exporting a report.

The risk an export carries that a screen does not: a spreadsheet outlives the
conversation and gets forwarded. So the two things worth pinning down are that
it contains the WHOLE result rather than the page the chat showed, and that it
contains only rows the person downloading it is allowed to see.
"""
import uuid
from decimal import Decimal
from io import BytesIO

import pytest
from jose import jwt
from openpyxl import load_workbook
from sqlalchemy.ext.asyncio import AsyncSession, async_sessionmaker

from app.core.config import settings
from app.models.po import PurchaseOrder
from app.models.vendor import Vendor
from app.services.controlled_query import MAX_ROWS_DETAIL

pytestmark = pytest.mark.asyncio

EXPORT = "/api/v1/assistant/export"


def _user_id(client) -> uuid.UUID:
    token = client.headers["Authorization"].split()[1]
    return uuid.UUID(jwt.decode(token, settings.JWT_SECRET_KEY,
                                algorithms=[settings.JWT_ALGORITHM])["sub"])


async def _seed_vendor(test_engine) -> Vendor:
    factory = async_sessionmaker(test_engine, class_=AsyncSession, expire_on_commit=False)
    async with factory() as db:
        v = Vendor(code=f"V-{uuid.uuid4().hex[:6]}", name="Export Vendor",
                   category="Services", contact_name="E", contact_email="e@t.test")
        db.add(v)
        await db.commit()
        await db.refresh(v)
        return v


async def _seed_pos(test_engine, vendor, creator, *, prefix, count, total="10.00"):
    factory = async_sessionmaker(test_engine, class_=AsyncSession, expire_on_commit=False)
    async with factory() as db:
        for i in range(count):
            db.add(PurchaseOrder(
                number=f"{prefix}-{i:04d}", title=f"{prefix}-{i:04d}", type=1,
                status="approved", subtotal=Decimal(total), total=Decimal(total),
                vendor_id=vendor.id, vendor_name=vendor.name, created_by=creator))
        await db.commit()


def _sheet(resp):
    return load_workbook(BytesIO(resp.content)).active


def _cells(ws):
    return [[c.value for c in row] for row in ws.iter_rows()]


async def test_the_export_is_a_real_workbook(test_engine, admin_client):
    vendor = await _seed_vendor(test_engine)
    tag = uuid.uuid4().hex[:6]
    await _seed_pos(test_engine, vendor, _user_id(admin_client), prefix=f"PO-X{tag}", count=3)

    r = await admin_client.post(EXPORT, json={
        "entity": "purchase_order", "question": "my orders",
        "select": ["number", "total"],
        "where": [{"field": "number", "op": "like", "value": f"PO-X{tag}"}],
    })
    assert r.status_code == 200
    assert "spreadsheetml" in r.headers["content-type"]
    assert ".xlsx" in r.headers["content-disposition"]
    # The name has to survive a cross-origin download or the browser invents one.
    assert "Content-Disposition" in r.headers.get("access-control-expose-headers", "")

    ws = _sheet(r)
    flat = [str(v) for row in _cells(ws) for v in row if v is not None]
    assert any("my orders" in v for v in flat), "the question heads the sheet"
    assert any(f"PO-X{tag}-0000" in v for v in flat)


async def test_the_export_ignores_the_chat_row_cap(test_engine, admin_client):
    """The whole point. A chat reply is capped at 100 rows; a report someone
    asked to have in full must not quietly stop there."""
    vendor = await _seed_vendor(test_engine)
    tag = uuid.uuid4().hex[:6]
    n = MAX_ROWS_DETAIL + 25
    await _seed_pos(test_engine, vendor, _user_id(admin_client), prefix=f"PO-B{tag}", count=n)

    r = await admin_client.post(EXPORT, json={
        "entity": "purchase_order", "select": ["number"],
        "where": [{"field": "number", "op": "like", "value": f"PO-B{tag}"}],
        # Exactly what a chat plan would have carried.
        "limit": 5,
    })
    ws = _sheet(r)
    numbers = {v for row in _cells(ws) for v in row
               if isinstance(v, str) and v.startswith(f"PO-B{tag}")}
    assert len(numbers) == n, f"expected all {n} rows, got {len(numbers)}"


async def test_money_lands_as_numbers_not_text(test_engine, admin_client):
    """Decimals travel as strings for precision; a column of text in a
    spreadsheet cannot be summed, which is the first thing anyone will try."""
    vendor = await _seed_vendor(test_engine)
    tag = uuid.uuid4().hex[:6]
    await _seed_pos(test_engine, vendor, _user_id(admin_client),
                    prefix=f"PO-M{tag}", count=2, total="12.34")

    r = await admin_client.post(EXPORT, json={
        "entity": "purchase_order", "select": ["number", "total"],
        "where": [{"field": "number", "op": "like", "value": f"PO-M{tag}"}],
    })
    values = [v for row in _cells(_sheet(r)) for v in row]
    assert any(isinstance(v, (int, float)) and abs(float(v) - 12.34) < 0.001
               for v in values), "money should be numeric in the sheet"


async def test_columns_are_headed_with_human_labels(test_engine, admin_client):
    vendor = await _seed_vendor(test_engine)
    tag = uuid.uuid4().hex[:6]
    await _seed_pos(test_engine, vendor, _user_id(admin_client), prefix=f"PO-H{tag}", count=1)

    r = await admin_client.post(EXPORT, json={
        "entity": "purchase_order", "select": ["number", "vendor_name"],
        "where": [{"field": "number", "op": "like", "value": f"PO-H{tag}"}],
    })
    flat = [str(v) for row in _cells(_sheet(r)) for v in row if v is not None]
    assert any("PO number" in v for v in flat), "not the raw field key"
    assert any("Vendor name" in v for v in flat)


async def test_a_grouped_export_carries_the_sql_total(test_engine, admin_client):
    vendor = await _seed_vendor(test_engine)
    tag = uuid.uuid4().hex[:6]
    await _seed_pos(test_engine, vendor, _user_id(admin_client),
                    prefix=f"PO-G{tag}", count=4, total="25.00")

    r = await admin_client.post(EXPORT, json={
        "entity": "purchase_order", "group_by": ["number"], "metrics": ["amount"],
        "where": [{"field": "number", "op": "like", "value": f"PO-G{tag}"}],
    })
    values = [v for row in _cells(_sheet(r)) for v in row]
    assert any(isinstance(v, (int, float)) and abs(float(v) - 100.0) < 0.001
               for v in values), "4 × 25.00 should appear as the total"
    flat = [str(v) for v in values if v is not None]
    assert any("Total" in v for v in flat)


# ── the property that matters most ────────────────────────────────────────────


async def test_the_export_is_scoped_to_whoever_downloads_it(
    test_engine, admin_client, requester_client
):
    """A spreadsheet is the easiest thing in the world to forward, so the rows
    in it must be the downloader's, not those of whoever ran the query first."""
    vendor = await _seed_vendor(test_engine)
    tag = uuid.uuid4().hex[:6]
    await _seed_pos(test_engine, vendor, _user_id(admin_client), prefix=f"PO-S{tag}A", count=3)
    await _seed_pos(test_engine, vendor, _user_id(requester_client), prefix=f"PO-S{tag}B", count=2)

    body = {"entity": "purchase_order", "select": ["number"],
            "where": [{"field": "number", "op": "like", "value": f"PO-S{tag}"}]}

    admin_nums = {v for row in _cells(_sheet(await admin_client.post(EXPORT, json=body)))
                  for v in row if isinstance(v, str) and v.startswith(f"PO-S{tag}")}
    req_nums = {v for row in _cells(_sheet(await requester_client.post(EXPORT, json=body)))
                for v in row if isinstance(v, str) and v.startswith(f"PO-S{tag}")}

    assert len(admin_nums) == 5, "unrestricted sees both sets"
    assert len(req_nums) == 2, "the requester sees only their own"
    assert not any(n.startswith(f"PO-S{tag}A") for n in req_nums)


async def test_an_unknown_entity_is_refused(admin_client):
    r = await admin_client.post(EXPORT, json={"entity": "salaries"})
    assert r.status_code == 422
