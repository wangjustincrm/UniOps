"""Server-side XLSX builder for the Budget Dashboard export.

Pure function: takes already-fetched data + returns .xlsx bytes. Only depends on
openpyxl + stdlib (no DB, no app-internal imports), so it is unit-testable without
a database.

Layout: one row per budget account, and each month occupies two columns — Plan
and NC posted actual — so a month reads left-to-right instead of plan/actual
living on separate rows. Year Plan / Year NC close the row, and a Grand Total row
closes the sheet.

With `partners` (the dashboard's "Expand vendor" toggle) a Vendor and a Vendor
Source column are added after Account Name, and every account row is followed by
one row per vendor carrying that vendor's NC actual — the same breakdown the
dashboard shows when an account is expanded. Plan is per account, never per
vendor, so vendor rows leave the Plan columns empty. The vendor rows are an Excel
outline level under their account (collapsible) and repeat the account's
category/code/name so a filter on Account Code keeps them. They do not count
toward the Grand Total: the account row above them already holds their sum.
"""
from __future__ import annotations

import io
from decimal import Decimal
from typing import Any

import openpyxl
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

_MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun",
           "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
_MONEY_FMT = "#,##0.00"
_LABELS = ("Category", "Account Code", "Account Name")
_VENDOR_LABELS = ("Vendor", "Vendor Source")
_GROUP_HEADER_ROW = 4                       # "Jan" spanning two columns
_SUB_HEADER_ROW = 5                         # "Plan" / "NC"

# Why a breakdown row has no vendor name — the dashboard's NO_VENDOR_LABEL.
_NO_VENDOR_LABEL = {
    "multi": "(several parties on the voucher)",
    "none": "(no vendor — internal entry)",
}
# Where a named vendor came from — the dashboard's "from voucher" badge.
_SOURCE_LABEL = {"line": "line", "voucher": "from voucher", "mixed": "partly from voucher"}


def _num(v: Any) -> float:
    """Coerce a money value (Decimal, str, int, None) to float. '' / None -> 0.0."""
    if v is None or v == "":
        return 0.0
    return float(Decimal(str(v)))


def _month(d: dict, m: int) -> float:
    """Read month `m` from a dict keyed by int (in-process) OR str (JSON)."""
    if m in d:
        return _num(d[m])
    return _num(d.get(str(m)))


def _interleave(plan: list[float | None], nc: list[float]) -> list[float | None]:
    """[p1,p2,…], [n1,n2,…] -> [p1,n1,p2,n2,…] — the month-pair column order."""
    out: list[float | None] = []
    for p, n in zip(plan, nc, strict=True):
        out.append(p)
        out.append(n)
    return out


def build_budget_actual_xlsx(*, accounts: list[dict[str, Any]],
                             nc_monthly: dict[str, dict],
                             fiscal_year: int, cost_center_label: str,
                             include_unposted: bool = False,
                             partners: dict[str, list[dict]] | None = None) -> bytes:
    """`partners` = crud.nc_partner_monthly_all's result ({account_id: [row…]});
    None exports one row per account, as before."""
    labels = _LABELS + (_VENDOR_LABELS if partners is not None else ())
    money_start = len(labels) + 1
    total_cols = len(labels) + 2 * (len(_MONTHS) + 1)   # +1 for the Year pair
    blank_labels = [""] * (len(labels) - len(_LABELS))

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Budget vs Actual"
    bold = Font(bold=True)
    centre = Alignment(horizontal="center", vertical="center")

    ws.append(["Budget vs Actual"])
    ws["A1"].font = bold
    ws.append([f"FY {fiscal_year} · Cost Center: {cost_center_label} · NC = "
               + ("posted + unposted (not yet tallied)" if include_unposted else "posted")])
    ws.append([])

    # Two header rows: month name spanning its Plan/NC pair, then the pair labels.
    for col, label in enumerate(labels, start=1):
        cell = ws.cell(row=_GROUP_HEADER_ROW, column=col, value=label)
        cell.font = bold
        cell.alignment = centre
        ws.merge_cells(start_row=_GROUP_HEADER_ROW, start_column=col,
                       end_row=_SUB_HEADER_ROW, end_column=col)

    col = money_start
    for label in (*_MONTHS, "Year"):
        group = ws.cell(row=_GROUP_HEADER_ROW, column=col, value=label)
        group.font = bold
        group.alignment = centre
        ws.merge_cells(start_row=_GROUP_HEADER_ROW, start_column=col,
                       end_row=_GROUP_HEADER_ROW, end_column=col + 1)
        for offset, sub in enumerate(("Plan", "NC")):
            cell = ws.cell(row=_SUB_HEADER_ROW, column=col + offset, value=sub)
            cell.font = bold
            cell.alignment = centre
        col += 2

    gt_plan = [0.0] * 12
    gt_nc = [0.0] * 12
    gt_plan_year = 0.0
    gt_nc_year = 0.0

    def money_row(values: list, *, strong: bool = False, level: int = 0) -> None:
        ws.append(values)
        r = ws.max_row
        for c in range(money_start, len(values) + 1):
            ws.cell(row=r, column=c).number_format = _MONEY_FMT
        if strong:
            for cell in ws[r]:
                cell.font = bold
        if level:
            ws.row_dimensions[r].outline_level = level

    ordered = sorted(accounts, key=lambda a: (a.get("l1_code", ""), a.get("account_code", "")))
    for a in ordered:
        aid = str(a["account_id"])
        plan_bm = a.get("plan_by_month") or {}
        nc_bm = nc_monthly.get(aid) or {}

        plan_months = [_month(plan_bm, m) for m in range(1, 13)]
        plan_year = _num(a.get("plan_year"))
        nc_months = [_month(nc_bm, m) for m in range(1, 13)]
        nc_year = sum(nc_months)

        acct_labels = [a.get("l1_code", ""), a.get("account_code", ""), a.get("account_name", "")]
        money_row([*acct_labels, *blank_labels,
                   *_interleave(plan_months, nc_months), plan_year, nc_year],
                  strong=partners is not None)

        for p in (partners or {}).get(aid, []):
            name = p.get("partner_name")
            source = p.get("source", "")
            p_months = [_month(p.get("by_month") or {}, m) for m in range(1, 13)]
            money_row([*acct_labels,
                       name or _NO_VENDOR_LABEL.get(source, "(no vendor)"),
                       _SOURCE_LABEL.get(source, "") if name else "",
                       *_interleave([None] * 12, p_months), None, sum(p_months)],
                      level=1)

        for i in range(12):
            gt_plan[i] += plan_months[i]
            gt_nc[i] += nc_months[i]
        gt_plan_year += plan_year
        gt_nc_year += nc_year

    ws.append([])
    money_row(["Grand Total", "", "", *blank_labels, *_interleave(gt_plan, gt_nc), gt_plan_year, gt_nc_year],
              strong=True)

    ws.freeze_panes = ws.cell(row=_SUB_HEADER_ROW + 1, column=money_start)
    ws.column_dimensions["A"].width = 12
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 34
    if partners is not None:
        ws.column_dimensions["D"].width = 36
        ws.column_dimensions["E"].width = 18
        ws.sheet_properties.outlinePr.summaryBelow = False   # account row sits above its vendors
    for c in range(money_start, total_cols + 1):
        ws.column_dimensions[get_column_letter(c)].width = 13

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
