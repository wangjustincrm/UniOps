"""The reconciliation report — the document the auditors get.

Deliberately the same shape as the QuickBooks report finance hands over today
(RBC CAD 2026.7RECON.pdf), because that is what the auditors already accept:

    Summary
      Statement beginning balance
      Cheques and payments cleared (N)
      Deposits and other credits cleared (N)
      Statement ending balance
      Register balance as of <date>
    Details
      Cheques and payments cleared (N)   date / type / ref / payee / amount
      Deposits and other credits cleared (N)

With one addition. QuickBooks prints no Outstanding section for July because
July happened to have none — register balance equalled statement closing. A
reconciliation that cannot show its uncleared items is only useful in the month
where there are none, so both sides of outstanding are always printed, and say
"none" when that is the answer.

The report is built from the SNAPSHOT once a period is finalized, never from a
re-query. A signed-off reconciliation has to keep saying what it said.
"""
import io
import re
from datetime import date
from decimal import Decimal
from xml.sax.saxutils import escape

ZERO = Decimal("0")

# NC's bank names, voucher summaries and payee names are Chinese. Helvetica has
# no Chinese glyphs, so every one of them printed as a row of black boxes.
_CJK_RE = re.compile(
    "[\u2e80-\u33ff\u3400-\u4dbf\u4e00-\u9fff\uf900-\ufaff"
    "\ufe30-\ufe4f\uff00-\uffef]+"
)

# Preferred first because a TTF is *embedded*: an auditor opening the report on
# a machine with no Chinese font installed still sees the bank's name. The CID
# font after it is only referenced by name, so it depends on the reader having
# one — better than boxes, worse than embedding. Helvetica last: the report
# still builds, the Chinese is still boxes, and nothing else is lost.
_CJK_TTF_CANDIDATES = (
    ("wqy-zenhei", "/usr/share/fonts/truetype/wqy/wqy-zenhei.ttc", 0),
    ("wqy-microhei", "/usr/share/fonts/truetype/wqy/wqy-microhei.ttc", 0),
    ("droid-fallback", "/usr/share/fonts/truetype/droid/DroidSansFallbackFull.ttf", 0),
)
_cjk_font_name = None


def _cjk_font() -> str:
    """Name of a registered font that can draw Chinese. Resolved once."""
    global _cjk_font_name
    if _cjk_font_name is not None:
        return _cjk_font_name
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.pdfbase.ttfonts import TTFont

    for name, path, index in _CJK_TTF_CANDIDATES:
        try:
            pdfmetrics.registerFont(TTFont(name, path, subfontIndex=index))
            _cjk_font_name = name
            return name
        except Exception:
            continue
    try:
        pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
        _cjk_font_name = "STSong-Light"
    except Exception:
        _cjk_font_name = "Helvetica"
    return _cjk_font_name


def _rich(text) -> str:
    """Paragraph markup: XML-escaped, with every Chinese run switched to the CJK
    font so the surrounding English keeps Helvetica (and its bold)."""
    t = escape(str(text if text is not None else ""))
    if not _CJK_RE.search(t):
        return t
    font = _cjk_font()
    if font == "Helvetica":
        return t
    return _CJK_RE.sub(lambda m: f'<font name="{font}">{m.group(0)}</font>', t)


def _cell(text, style):
    """A table cell. Left as a plain string unless it has Chinese in it — that
    keeps the all-English case laid out exactly as it was."""
    from reportlab.platypus import Paragraph

    t = str(text if text is not None else "")
    return Paragraph(_rich(t), style) if _CJK_RE.search(t) else t


def _money(v) -> str:
    d = Decimal(str(v or 0))
    return f"{d:,.2f}"


def build_snapshot(rec, account_label: str, summary: dict, cleared_debits: list,
                   cleared_credits: list, outstanding_bank: list,
                   outstanding_book: list, prepared_by: str) -> dict:
    """Everything the report needs, frozen. Plain JSON — no ORM objects — so it
    survives a schema change as well as an NC re-sync."""
    return {
        "account": account_label,
        "period_start": rec.period_start.isoformat(),
        "period_end": rec.period_end.isoformat(),
        "currency": rec.currency,
        "prepared_by": prepared_by,
        "summary": summary,
        "cleared_payments": cleared_debits,
        "cleared_deposits": cleared_credits,
        "outstanding_bank": outstanding_bank,
        "outstanding_book": outstanding_book,
    }


def _base_only_note(snapshot: dict) -> str | None:
    """On a foreign-currency account, the ledger lines left out because they have
    no amount in that currency. An auditor comparing to NC's CAD 科目余额表 will
    see a different register balance; this line is the reason, with the number."""
    s = snapshot.get("summary") or {}
    n = int(s.get("base_only_count") or 0)
    if not n:
        return None
    return (f"Register balance is in {snapshot['currency']} (NC original currency). "
            f"{n} ledger line(s) with no {snapshot['currency']} amount — FX "
            f"revaluation / CAD-only adjustments, CAD "
            f"{_money(s.get('base_only_local_total'))} — are not bank movements "
            f"and are excluded.")


# ── PDF ────────────────────────────────────────────────────────────────────────

def render_pdf(snapshot: dict, reconciled_on: date, reconciled_by: str) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import letter, landscape
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import (
        PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle,
    )

    styles = getSampleStyleSheet()
    # ReportLab's default leading is 12 regardless of font size, so anything
    # larger than ~10pt overlaps the line below it unless leading is set too.
    h1 = ParagraphStyle("h1", parent=styles["Normal"], fontName="Helvetica-Bold",
                        fontSize=13, leading=17)
    h2 = ParagraphStyle("h2", parent=styles["Normal"], fontName="Helvetica-Bold",
                        fontSize=10.5, leading=14, spaceBefore=10, spaceAfter=4)
    small = ParagraphStyle("small", parent=styles["Normal"], fontSize=8.5, leading=11,
                           textColor=colors.HexColor("#5D6B6A"))
    # Only used for cells that actually contain Chinese; matches the FONTSIZE the
    # table style applies to the plain-string cells around it.
    cell = ParagraphStyle("cell", parent=styles["Normal"], fontSize=7.5, leading=9.5)

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(letter),
                            leftMargin=14 * mm, rightMargin=14 * mm,
                            topMargin=12 * mm, bottomMargin=12 * mm,
                            title=f"Bank Reconciliation {snapshot['account']} "
                                  f"{snapshot['period_end']}")
    flow = []
    s = snapshot["summary"]
    flow.append(Paragraph(_rich(f"{snapshot['account']}, Period Ending "
                                f"{snapshot['period_end']}"), h1))
    flow.append(Paragraph("RECONCILIATION REPORT", h2))
    flow.append(Paragraph(
        f"Reconciled on: {reconciled_on.isoformat()} &nbsp;&nbsp; "
        f"Reconciled by: {_rich(reconciled_by)}", small))
    flow.append(Paragraph(
        "Any changes made to transactions after this date are not included in this "
        "report.", small))
    flow.append(Spacer(1, 8))

    n_pay, n_dep = len(snapshot["cleared_payments"]), len(snapshot["cleared_deposits"])
    summary_rows = [
        ["Summary", snapshot["currency"]],
        ["Statement beginning balance", _money(s.get("statement_opening"))],
        [f"Cheques and payments cleared ({n_pay})",
         "-" + _money(s.get("cleared_debit_total"))],
        [f"Deposits and other credits cleared ({n_dep})",
         _money(s.get("cleared_credit_total"))],
        ["Statement ending balance", _money(s.get("statement_closing"))],
        [f"Register balance as of {snapshot['period_end']}", _money(s.get("book_closing"))],
        ["Difference", _money(s.get("difference"))],
    ]
    t = Table(summary_rows, colWidths=[150 * mm, 40 * mm])
    t.setStyle(TableStyle([
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("ALIGN", (1, 0), (1, -1), "RIGHT"),
        ("LINEBELOW", (0, 0), (-1, 0), 0.6, colors.black),
        ("LINEABOVE", (0, -1), (-1, -1), 0.4, colors.grey),
        ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
    ]))
    flow.append(t)
    if note := _base_only_note(snapshot):
        flow.append(Spacer(1, 4))
        flow.append(Paragraph(escape(note), small))
    flow.append(Spacer(1, 10))

    def detail(title: str, rows: list, negative: bool):
        flow.append(Paragraph(f"{title} ({len(rows)})", h2))
        if not rows:
            flow.append(Paragraph("None.", small))
            return
        data = [["DATE", "TYPE", "REF NO.", "PAYEE / DESCRIPTION",
                 f"AMOUNT ({snapshot['currency']})"]]
        for r in rows:
            amt = _money(r.get("amount"))
            data.append([r.get("date", ""),
                         _cell(r.get("type", ""), cell),
                         _cell(r.get("ref", ""), cell),
                         _cell((r.get("payee") or "")[:70], cell),
                         ("-" + amt) if negative else amt])
        data.append(["", "", "", "Total",
                     ("-" if negative else "") + _money(
                         sum(Decimal(str(r.get("amount") or 0)) for r in rows))])
        tt = Table(data, colWidths=[22 * mm, 30 * mm, 30 * mm, 148 * mm, 30 * mm],
                   repeatRows=1)
        tt.setStyle(TableStyle([
            ("FONTSIZE", (0, 0), (-1, -1), 7.5),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#5D6B6A")),
            ("ALIGN", (4, 0), (4, -1), "RIGHT"),
            ("LINEBELOW", (0, 0), (-1, 0), 0.5, colors.black),
            ("LINEABOVE", (0, -1), (-1, -1), 0.5, colors.black),
            ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
            ("ROWBACKGROUNDS", (0, 1), (-1, -2),
             [colors.white, colors.HexColor("#F5F7F7")]),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
            ("TOPPADDING", (0, 0), (-1, -1), 2),
        ]))
        flow.append(tt)

    detail("Cheques and payments cleared", snapshot["cleared_payments"], negative=True)
    flow.append(PageBreak())
    detail("Deposits and other credits cleared", snapshot["cleared_deposits"],
           negative=False)
    detail("Outstanding statement lines — on the bank, not in the ledger",
           snapshot["outstanding_bank"], negative=False)
    detail("Outstanding ledger lines — in the ledger, not on the bank",
           snapshot["outstanding_book"], negative=False)

    doc.build(flow)
    return buf.getvalue()


# ── Excel ──────────────────────────────────────────────────────────────────────

def render_xlsx(snapshot: dict, reconciled_on: date, reconciled_by: str) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    bold = Font(bold=True)
    head = PatternFill("solid", fgColor="E4EFEC")
    money = "#,##0.00"

    ws = wb.active
    ws.title = "Summary"
    s = snapshot["summary"]
    ws.append([f"{snapshot['account']} — Reconciliation Report"])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([f"Period {snapshot['period_start']} to {snapshot['period_end']}"])
    ws.append([f"Reconciled on {reconciled_on.isoformat()} by {reconciled_by}"])
    ws.append([])
    for label, value in (
        ("Statement beginning balance", s.get("statement_opening")),
        (f"Cheques and payments cleared ({len(snapshot['cleared_payments'])})",
         Decimal(str(s.get("cleared_debit_total") or 0)) * -1),
        (f"Deposits and other credits cleared ({len(snapshot['cleared_deposits'])})",
         s.get("cleared_credit_total")),
        ("Statement ending balance", s.get("statement_closing")),
        ("Register balance", s.get("book_closing")),
        ("Difference", s.get("difference")),
    ):
        ws.append([label, float(Decimal(str(value or 0)))])
        ws.cell(ws.max_row, 2).number_format = money
    ws.cell(ws.max_row, 1).font = bold
    ws.cell(ws.max_row, 2).font = bold
    if note := _base_only_note(snapshot):
        ws.append([])
        ws.append([note])
    ws.column_dimensions["A"].width = 46
    ws.column_dimensions["B"].width = 18

    for title, rows, sign in (
        ("Cleared payments", snapshot["cleared_payments"], -1),
        ("Cleared deposits", snapshot["cleared_deposits"], 1),
        ("Outstanding — bank", snapshot["outstanding_bank"], 1),
        ("Outstanding — ledger", snapshot["outstanding_book"], 1),
    ):
        sh = wb.create_sheet(title[:31])
        sh.append(["Date", "Type", "Ref No.", "Payee / description",
                   f"Amount ({snapshot['currency']})"])
        for c in sh[1]:
            c.font = bold
            c.fill = head
        for r in rows:
            sh.append([r.get("date", ""), r.get("type", ""), r.get("ref", ""),
                       r.get("payee", ""),
                       float(Decimal(str(r.get("amount") or 0)) * sign)])
            sh.cell(sh.max_row, 5).number_format = money
        for col, width in zip("ABCDE", (12, 18, 18, 62, 16)):
            sh.column_dimensions[col].width = width
        sh.freeze_panes = "A2"
        sh.auto_filter.ref = f"A1:E{max(sh.max_row, 1)}"
        for c in sh["D"]:
            c.alignment = Alignment(wrap_text=False)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
