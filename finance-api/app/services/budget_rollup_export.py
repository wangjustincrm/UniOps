"""XLSX for the Budget-vs-Actual roll-up — the report finance sends onward.

Pure function: takes the roll-up payload, returns bytes. No DB, no app imports
beyond openpyxl, so it is unit-testable without a database (same shape as
predreal_export.py).

Exports the tree ALREADY EXPANDED: every cost centre under every group, whether
or not it was expanded on screen. A collapsed row on a page is a convenience; a
collapsed row in a spreadsheet is missing data (user, 2026-09-18).
"""
from __future__ import annotations

import io
from decimal import Decimal
from typing import Any

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

_MONEY = "#,##0.00"
_PCT = "0.0"

# Excel is stricter than openpyxl about three things, and says only "we found a
# problem with some content" when any of them is off — it does not name the
# part, and the file still opens in openpyxl, so none of this shows up in a
# round-trip test. Hence the explicit helpers and the assertions in
# tests/test_budget_rollup_export.py.
#
#  1. Colours are aRGB. A six-digit value gets an implicit "00" alpha, i.e.
#     fully transparent, which is not what any of these mean.
#  2. A solid patternFill needs BOTH fgColor and bgColor.
#  3. A border element is expected to carry all of its sides, not only the one
#     being drawn.


def _rgb(hex6: str) -> str:
    return f"FF{hex6}"


def _fill(hex6: str) -> PatternFill:
    colour = _rgb(hex6)
    return PatternFill(fill_type="solid", start_color=colour, end_color=colour)


def _border(*, top: Side | None = None, bottom: Side | None = None) -> Border:
    blank = Side()
    return Border(left=blank, right=blank, diagonal=blank,
                  top=top or blank, bottom=bottom or blank)


_HEAD_FILL = _fill("F5F5F5")
_GROUP_FILL = _fill("FAFAFA")
_POLICY_FILL = _fill("FFFBEB")
_TOTAL_FILL = _fill("EDF2F7")
_THIN = Side(style="thin", color=_rgb("D4D4D4"))
_MED = Side(style="medium", color=_rgb("A3A3A3"))


def _num(v: Any) -> float | None:
    """Money/percent to float. None stays None so the cell is blank rather than
    0.00 — "no plan" and "a plan of zero" are different answers."""
    if v is None or v == "":
        return None
    return float(Decimal(str(v)))


_MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun",
           "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]


def build_rollup_xlsx(*, rollup: dict[str, Any], group_by: str,
                      window_label: str, by_month: bool = False) -> bytes:
    """`by_month` puts a Plan / Actual pair for every month of the window in
    front of the window totals — the same expansion the page shows. No monthly
    variance: finance asked for the two figures only (user, 2026-09-25)."""
    by_department = group_by == "department"
    nodes = rollup["by_department" if by_department else "by_expense_centre"]
    level_header = "Department" if by_department else "Expense centre"

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Budget vs Actual"

    year = rollup["fiscal_year"]
    ws["A1"] = f"Budget vs Actual — by {level_header.lower()}"
    ws["A1"].font = Font(size=14, bold=True)
    ws["A2"] = (f"FY {year} · {window_label} · amounts in CAD · "
                f"actual = NC posted")
    ws["A2"].font = Font(size=10, color=_rgb("737373"))

    months = (list(range(rollup["month_from"], rollup["month_to"] + 1))
              if by_month else [])
    # Column layout: label, code, [month plan/actual pairs], the eight measures.
    first_measure = 3 + 2 * len(months)
    pct_cols = {first_measure + 3, first_measure + 7}

    headers = [
        level_header, "Code",
        *[h for m in months for h in (f"{_MONTHS[m - 1]} Plan", f"{_MONTHS[m - 1]} Actual")],
        f"Plan {window_label}", f"Actual {window_label}", "Variance", "Variance %",
        "Full-year plan", "Actual YTD", "Remaining", "Consumed %",
    ]
    row = 4
    for col, h in enumerate(headers, 1):
        c = ws.cell(row=row, column=col, value=h)
        c.font = Font(bold=True, size=10)
        c.fill = _HEAD_FILL
        c.border = _border(bottom=_MED)
        c.alignment = Alignment(horizontal="left" if col <= 2 else "right",
                                vertical="center", wrap_text=True)
    ws.freeze_panes = ws.cell(row=row + 1, column=1)

    def write(label: str, code: str | None, m: dict[str, Any], *,
              indent: int = 0, bold: bool = False, fill: PatternFill | None = None,
              top_border: bool = False) -> None:
        nonlocal row
        row += 1
        per_month = {e["month"]: e for e in m.get("monthly") or []}
        cells = [(1, label), (2, code or "")]
        for i, month in enumerate(months):
            e = per_month.get(month, {})
            cells += [(3 + 2 * i, _num(e.get("plan"))), (4 + 2 * i, _num(e.get("actual")))]
        cells += [
            (first_measure + i, _num(m.get(key))) for i, key in enumerate((
                "plan_period", "actual_period", "variance_period", "variance_period_pct",
                "plan_full_year", "actual_ytd", "remaining_full_year", "consumed_pct"))
        ]
        for col, value in cells:
            c = ws.cell(row=row, column=col, value=value)
            if col == 1:
                # `indent` is only valid alongside a horizontal alignment —
                # Excel reports "we found a problem with some content" on an
                # <alignment indent="2"/> with no horizontal attribute, and
                # then silently drops the styling it recovered.
                c.alignment = Alignment(horizontal="left", indent=indent)
            elif col >= 3:
                c.number_format = _PCT if col in pct_cols else _MONEY
            if bold:
                c.font = Font(bold=True)
            if fill is not None:
                c.fill = fill
            c.border = _border(top=_MED if top_border else _THIN)

    write("Whole company", "", rollup["company"], bold=True, fill=_TOTAL_FILL)

    for node in nodes:
        write(node["label"], node["code"], node, bold=True, fill=_GROUP_FILL)
        # Always written out, never only when expanded on screen.
        for child in node["children"]:
            write(child["cost_center_name"], child["cost_center_code"], child, indent=2)

    unalloc = rollup["unallocated"]
    write("Category-level (not allocated)", "",
          {"actual_period": unalloc["actual_period"], "actual_ytd": unalloc["actual_ytd"],
           "monthly": unalloc.get("monthly")},
          bold=True, fill=_POLICY_FILL, top_border=True)
    for b in unalloc["breakdown"]:
        write(b["label"], "",
              {"actual_period": b["actual_period"], "actual_ytd": b["actual_ytd"],
               "monthly": b.get("monthly")},
              indent=2, fill=_POLICY_FILL)

    rec = rollup["reconciliation"]
    # The same proof the page carries: placed + unallocated == everything posted.
    write(f"Total actual {window_label}", "",
          {"actual_period": rec["total_actual_period"],
           "monthly": rec.get("total_monthly")},
          bold=True, fill=_TOTAL_FILL, top_border=True)
    row += 1
    # NOT "= x + y": a cell whose text starts with "=" is a formula to Excel,
    # and one that does not parse is silently dropped — the repair log reads
    # "Removed Records: Formula from /xl/worksheets/sheet1.xml".
    ws.cell(row=row, column=1,
            value=(f"of which {rec['placed_actual_period']} in cost centres "
                   f"and {rec['unallocated_actual_period']} category-level")
            ).font = Font(size=9, color=_rgb("737373"))

    widths = [34, 16, *[14] * (2 * len(months)), 16, 16, 14, 11, 16, 16, 14, 12]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
