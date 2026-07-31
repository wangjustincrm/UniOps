"""NC AP export (parallel-run) — batches, assembly, xlsx, API."""
import uuid
from datetime import date, datetime, timezone, timedelta
from decimal import Decimal

import pytest
import pytest_asyncio
from httpx import ASGITransport, AsyncClient
from jose import jwt
from sqlalchemy import select
from app.services.posting import emit_event

from app.core.config import settings as app_settings
from app.db.base import get_db
from app.main import app


def _h(role="finance_manager"):
    tok = jwt.encode({"sub": str(uuid.uuid4()), "role": role,
                      "exp": datetime.now(timezone.utc) + timedelta(hours=1)},
                     app_settings.jwt_secret_key, algorithm=app_settings.jwt_algorithm)
    return {"Authorization": f"Bearer {tok}"}


@pytest_asyncio.fixture
async def client(db_session):
    async def _override():
        yield db_session
    app.dependency_overrides[get_db] = _override
    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as c:
        yield c
    app.dependency_overrides.clear()


async def _mk_ap(db, *, source="epms", amount="100.00", tax="13.00",
                 number="AP-2026-0100", vendor="ACME Ltd", src_id=None,
                 currency="CAD"):
    from app.models.ap_invoice import ApInvoice
    inv = ApInvoice(ap_invoice_number=number, source=source,
                    source_invoice_id=src_id or uuid.uuid4(),
                    vendor_name=vendor, amount=Decimal(amount),
                    tax_amount=Decimal(tax),
                    total_amount=Decimal(amount) + Decimal(tax),
                    currency=currency, invoice_date=date(2026, 7, 10), status="posted")
    db.add(inv)
    await db.flush()
    return inv


async def _mk_accrual(db, ap, account="510102"):
    await emit_event(
        db, source_service="finance", source_doc_type="ap_invoice",
        source_doc_id=ap.id, source_doc_number=ap.ap_invoice_number,
        event_type="accrual", prepared_by=uuid.uuid4(),
        lines=[{"line_role": "purchase_expense", "account_code": account,
                "debit": ap.amount, "currency": "CAD"},
               {"line_role": "sales_tax", "account_code": "1180",
                "debit": ap.tax_amount, "tax_code": "HST_ON", "currency": "CAD"},
               {"line_role": "accounts_payable", "account_code": "220201",
                "credit": ap.total_amount, "currency": "CAD"}])


# ── classify_expense_account unit tests ────────────────────────────────────────

def test_classify_expense_account_moh():
    from app.services.nc_ap_export import classify_expense_account
    assert classify_expense_account("MOH-0106-E01", None) == "510101"
    assert classify_expense_account("MOH-0104-P01", "Production") == "510101"


def test_classify_expense_account_rd():
    from app.services.nc_ap_export import classify_expense_account
    assert classify_expense_account("RD-001", None) == "5301"


def test_classify_expense_account_sell():
    from app.services.nc_ap_export import classify_expense_account
    assert classify_expense_account("SELL-0107-S03", None) == "660101"


def test_classify_expense_account_ga():
    from app.services.nc_ap_export import classify_expense_account
    assert classify_expense_account("GA-0105", None) == "6602"
    assert classify_expense_account("GA-0101", "HR") == "6602"


def test_classify_expense_account_no_cc_engineering():
    from app.services.nc_ap_export import classify_expense_account
    assert classify_expense_account(None, "Engineering") == "510101"
    assert classify_expense_account(None, "Production") == "510101"


def test_classify_expense_account_no_cc_sales():
    from app.services.nc_ap_export import classify_expense_account
    assert classify_expense_account(None, "Sales") == "660101"
    assert classify_expense_account(None, "Marketing") == "660101"
    assert classify_expense_account(None, "BD") == "660101"
    assert classify_expense_account(None, "E-COM") == "660101"


def test_classify_expense_account_no_cc_rnd():
    from app.services.nc_ap_export import classify_expense_account
    assert classify_expense_account(None, "R&D") == "5301"


def test_classify_expense_account_no_cc_unknown():
    from app.services.nc_ap_export import classify_expense_account
    assert classify_expense_account(None, "Maintenance") == "6602"
    assert classify_expense_account(None, None) == "6602"
    assert classify_expense_account(None, "") == "6602"


# ── integration tests ──────────────────────────────────────────────────────────

async def test_dept_fallback_not_fooled_by_substring(db_session):
    """Department-name fallback must NOT match 'Sales' for 'After-Sales Service'."""
    from app.models.mirrors import Department, InvoicePoAllocation, PurchaseRequest
    from app.services.nc_ap_export import build_export_rows
    db_session.add(Department(id=uuid.uuid4(), code="0110", name="Sales", is_active=True))
    src, po = uuid.uuid4(), uuid.uuid4()
    db_session.add(PurchaseRequest(id=uuid.uuid4(), po_id=po, cost_center_id=None,
                                   budget_code=None, department_name="After-Sales Service",
                                   created_by=uuid.uuid4()))
    db_session.add(InvoicePoAllocation(id=uuid.uuid4(), invoice_id=src, po_id=po,
                                       allocated_amount=Decimal("10.00"),
                                       allocated_tax=Decimal("0.00")))
    await db_session.flush()
    ap = await _mk_ap(db_session, src_id=src, number="AP-2026-0300")
    await _mk_accrual(db_session, ap)
    heads, bodies, errors = await build_export_rows(db_session, [ap.id])
    assert errors == []
    assert bodies[0]["department"] == ""  # must NOT resolve to Sales (0110)

async def test_build_rows_epms_allocations(db_session):
    from app.models.mirrors import CostCenter, Department, InvoicePoAllocation, PurchaseRequest, BudgetAccount
    from app.services.nc_ap_export import build_export_rows
    src = uuid.uuid4()
    po1, po2 = uuid.uuid4(), uuid.uuid4()
    cc_id = uuid.uuid4()
    dept_id = uuid.uuid4()
    db_session.add(Department(id=dept_id, code="0104", name="Production", is_active=True))
    db_session.add(CostCenter(id=cc_id, code="MOH-0104-P01", name="Production CC",
                              department_id=dept_id))
    db_session.add(BudgetAccount(id=uuid.uuid4(), code="CRM004", name="Depreciation", is_active=True))
    db_session.add_all([
        PurchaseRequest(id=uuid.uuid4(), po_id=po1, cost_center_id=cc_id,
                        budget_code="CRM004", department_name="Production",
                        created_by=uuid.uuid4()),
        PurchaseRequest(id=uuid.uuid4(), po_id=po2, cost_center_id=None,
                        budget_code=None, department_name="Maintenance",
                        created_by=uuid.uuid4()),
        InvoicePoAllocation(id=uuid.uuid4(), invoice_id=src, po_id=po1,
                            allocated_amount=Decimal("60.00"), allocated_tax=Decimal("7.80")),
        InvoicePoAllocation(id=uuid.uuid4(), invoice_id=src, po_id=po2,
                            allocated_amount=Decimal("40.00"), allocated_tax=Decimal("5.20")),
    ])
    await db_session.flush()
    ap = await _mk_ap(db_session, src_id=src)
    await _mk_accrual(db_session, ap)

    heads, bodies, errors = await build_export_rows(db_session, [ap.id])
    assert errors == []
    assert len(heads) == 1 and len(bodies) == 2
    h = heads[0]
    assert h["seq"] == 0 and h["billno"] == "" and h["ap_number"] == "AP-2026-0100"
    # head department = first body row's dept CODE
    assert h["department"] == "0104"
    b1 = next(b for b in bodies if b["money"] == "67.80")
    # cost_center: MOH-0104-P01 → 'P01' via NC_CC_BY_UNIOPS
    assert b1["cost_center"] == "P01"
    # revexp is the budget CODE directly
    assert b1["revexp"] == "CRM004"
    # notax/tax always emitted as zero — NC recomputes from money + tax_rate
    assert b1["notax"] == "0.00" and b1["tax"] == "0.00" and b1["money"] == "67.80"
    # account_path: MOH prefix → 510101
    assert b1["account_path"] == "510101"
    # tax_rate still derived from the real amounts (7.80 / 60.00)
    assert b1["tax_code"] == "001" and b1["tax_rate"] == "13.00"
    # CAD → buysell '2'
    assert b1["buysell"] == "2"
    # pay_term code
    assert b1["pay_term"] == "FH01"
    # obj_type code
    assert b1["obj_type"] == "1"
    b2 = next(b for b in bodies if b["money"] == "45.20")
    # no CC, Maintenance dept → unresolvable dept code → ''
    assert b2["cost_center"] == "" and b2["revexp"] == ""
    # Maintenance → no keyword match → 6602
    assert b2["account_path"] == "6602"


async def test_build_rows_oa_pa_chain(db_session):
    from app.models.mirrors import CostCenter, Department, ExpenseInvoice, User, BudgetAccount
    from app.models.pa import PaymentApplication
    from app.services.nc_ap_export import build_export_rows
    src, pa_id, cc_id, creator = uuid.uuid4(), uuid.uuid4(), uuid.uuid4(), uuid.uuid4()
    dept_id = uuid.uuid4()
    db_session.add(Department(id=dept_id, code="0101", name="Administration", is_active=True))
    db_session.add(CostCenter(id=cc_id, code="GA-0101", name="Admin CC",
                              department_id=dept_id))
    db_session.add(BudgetAccount(id=uuid.uuid4(), code="CRM010", name="Office Supplies", is_active=True))
    db_session.add(User(id=creator, email="a@x.com", full_name="Alice Wong"))
    db_session.add(ExpenseInvoice(id=src, pa_id=pa_id))
    db_session.add(PaymentApplication(
        id=pa_id, pa_number="PA-1", title="t", pa_type="PA-DIR", status="paid",
        vendor_id=uuid.uuid4(), vendor_name="ACME Ltd",
        payment_amount=Decimal("113.00"), currency="CAD",
        cost_center_id=cc_id, budget_account_code="CRM010", created_by=creator))
    await db_session.flush()
    ap = await _mk_ap(db_session, source="oa", number="AP-2026-0101", src_id=src)
    await _mk_accrual(db_session, ap)

    heads, bodies, errors = await build_export_rows(db_session, [ap.id])
    assert errors == [] and len(bodies) == 1
    b = bodies[0]
    # GA-0101 → 'HR' via NC_CC_BY_UNIOPS
    assert b["cost_center"] == "HR"
    # revexp = budget CODE directly
    assert b["revexp"] == "CRM010"
    assert b["employee"] == "Alice Wong"
    assert heads[0]["employee"] == "Alice Wong"
    # GA prefix → account 6602
    assert b["account_path"] == "6602"
    # dept code via CC.department_id → dept code '0101'
    assert b["department"] == "0101"
    # buysell CAD → '2'
    assert b["buysell"] == "2"
    assert b["pay_term"] == "FH01"
    assert b["obj_type"] == "1"


async def test_build_rows_errors(db_session):
    from app.services.nc_ap_export import build_export_rows
    ap_draft = await _mk_ap(db_session, number="AP-2026-0102")
    ap_draft.status = "draft"
    ap_noacc = await _mk_ap(db_session, number="AP-2026-0103")
    await db_session.flush()
    heads, bodies, errors = await build_export_rows(db_session, [ap_draft.id, ap_noacc.id])
    assert heads == [] and bodies == []
    reasons = {e["ap_number"]: e["reason"] for e in errors}
    assert reasons["AP-2026-0102"] == "status draft"
    assert reasons["AP-2026-0103"] == "no accrual posting"


async def test_export_batch_and_ap_columns(db_session):
    from app.models.ap_invoice import ApInvoice
    from app.models.nc_export import NcExportBatch
    b = NcExportBatch(exported_by=uuid.uuid4(),
                      exported_at=datetime.now(timezone.utc), ap_count=2,
                      filename="NC-AP-x.xlsx")
    db_session.add(b)
    inv = ApInvoice(ap_invoice_number="AP-2026-0001", source="epms",
                    source_invoice_id=uuid.uuid4(), amount=Decimal("100"),
                    tax_amount=Decimal("13"), total_amount=Decimal("113"),
                    currency="CAD", invoice_date=date(2026, 7, 1), status="posted")
    db_session.add(inv)
    await db_session.flush()
    inv.nc_exported_at = datetime.now(timezone.utc)
    inv.nc_export_batch_id = b.id
    await db_session.flush()
    got = (await db_session.execute(select(ApInvoice).where(
        ApInvoice.id == inv.id))).scalar_one()
    assert got.nc_export_batch_id == b.id


async def test_new_mirror_subsets_readable(db_session):
    from app.models.mirrors import ExpenseInvoice, InvoicePoAllocation, PurchaseRequest
    po = uuid.uuid4()
    db_session.add(PurchaseRequest(id=uuid.uuid4(), po_id=po,
                                   cost_center_id=uuid.uuid4(), budget_code="CRM004",
                                   department_name="Engineering", created_by=uuid.uuid4()))
    db_session.add(InvoicePoAllocation(id=uuid.uuid4(), invoice_id=uuid.uuid4(),
                                       po_id=po, allocated_amount=Decimal("50"),
                                       allocated_tax=Decimal("6.50")))
    db_session.add(ExpenseInvoice(id=uuid.uuid4(), pa_id=uuid.uuid4()))
    await db_session.flush()
    got = (await db_session.execute(select(PurchaseRequest).where(
        PurchaseRequest.po_id == po))).scalar_one()
    assert got.department_name == "Engineering"


def test_write_xlsx_structure(tmp_path):
    import openpyxl
    from app.services.nc_ap_export import write_xlsx
    heads = [{"seq": 0, "billno": "", "ap_number": "AP-1", "ap_type": "Payable of Expense",
              "busi_process": "AP01", "billdate": "2026-07-10",
              "busidate": "2026-07-10", "obj_type": "1", "supplier": "ACME",
              "department": "0104", "employee": "", "revexp": "CRM004",
              "currency": "CAD", "ap_type_code": "F1-Cxx-017", "tax_country": "Canada"},
             {"seq": 1, "billno": "", "ap_number": "AP-2", "ap_type": "Payable of Expense",
              "busi_process": "AP01", "billdate": "2026-07-11",
              "busidate": "2026-07-11", "obj_type": "1", "supplier": "Beta",
              "department": "", "employee": "Alice Wong", "revexp": "",
              "currency": "CAD", "ap_type_code": "F1-Cxx-017", "tax_country": "Canada"}]
    body_base = {"account_path": "510101", "invoice_no": "INV-9",
                 "summary": "ACME PO-1", "pay_term": "FH01",
                 "obj_type": "1", "supplier": "ACME", "department": "0104",
                 "cost_center": "P01", "employee": "", "revexp": "CRM004",
                 "currency": "CAD", "rate": "1", "money": "67.80", "qty": "",
                 "tax_code": "001", "tax_rate": "13.00", "tax_price": "0.00000000",
                 "notax": "60.00", "tax": "7.80", "taxtype": "02",
                 "department2": "0104", "buysell": "2"}
    bodies = [dict(body_base, seq=0), dict(body_base, seq=0, notax="40.00"),
              dict(body_base, seq=1)]
    data = write_xlsx(heads, bodies)
    p = tmp_path / "out.xlsx"
    p.write_bytes(data)
    ws = openpyxl.load_workbook(p)["Sheet1"]
    assert str(ws["A2"].value).startswith('"payablebill_$head')   # tech row kept
    assert ws["A3"].value == "0" and ws["C3"].value in (None, "")     # head seq + billno (now empty)
    assert ws["A4"].value == "1" and ws["C4"].value in (None, "")
    assert ws.cell(5, 1).value in (None, "")                      # blank separator
    assert str(ws["A6"].value).startswith('"bodys')               # body tech row
    assert ws["A7"].value == "0" and ws["B7"].value == "510101"
    assert ws["A9"].value == "1"                                  # 3rd body row -> doc 1
    assert ws["U7"].value == "60.00"                              # notax col


async def test_export_endpoint_streams_and_marks(client, db_session):
    from app.models.nc_export import NcExportBatch
    ap = await _mk_ap(db_session, number="AP-2026-0200")
    await _mk_accrual(db_session, ap)
    r = await client.post("/finance/v1/ap/nc-export",
                          json={"ap_ids": [str(ap.id)]}, headers=_h())
    assert r.status_code == 200, r.text
    assert r.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert "AP-2026-0200.xlsx" in r.headers["content-disposition"]
    assert len(r.content) > 1000
    await db_session.refresh(ap)
    assert ap.nc_exported_at is not None and ap.nc_export_batch_id is not None
    batch = (await db_session.execute(select(NcExportBatch))).scalars().first()
    assert batch is not None and batch.ap_count == 1


async def test_export_endpoint_allows_ap_clerk(client, db_session):
    """AP clerk (any finance role) may run the NC export — it is a routine AP
    handoff to NC65, not chart-of-accounts management. Regression for the
    mis-gated finance.coa.manage lock that 403'd every ap_clerk."""
    ap = await _mk_ap(db_session, number="AP-2026-0210")
    await _mk_accrual(db_session, ap)
    r = await client.post("/finance/v1/ap/nc-export",
                          json={"ap_ids": [str(ap.id)]}, headers=_h(role="ap_clerk"))
    assert r.status_code == 200, r.text


async def test_export_endpoint_guards(client, db_session):
    ap = await _mk_ap(db_session, number="AP-2026-0201")   # no accrual
    r = await client.post("/finance/v1/ap/nc-export",
                          json={"ap_ids": [str(ap.id)]}, headers=_h())
    assert r.status_code == 409
    assert r.json()["detail"]["errors"][0]["reason"] == "no accrual posting"
    r403 = await client.post("/finance/v1/ap/nc-export",
                             json={"ap_ids": [str(ap.id)]}, headers=_h(role="requester"))
    assert r403.status_code == 403
    rb = await client.get("/finance/v1/ap/nc-export/batches", headers=_h())
    assert rb.status_code == 200 and rb.json() == []


async def test_buysell_non_cad(db_session):
    """Non-CAD invoice → buysell '4'."""
    from app.services.nc_ap_export import build_export_rows
    ap = await _mk_ap(db_session, number="AP-2026-0300", currency="USD")
    await _mk_accrual(db_session, ap)
    _, bodies, errors = await build_export_rows(db_session, [ap.id])
    assert errors == []
    assert bodies[0]["buysell"] == "4"


async def test_ap_list_pagination_and_search(client, db_session):
    for i in range(3):
        await _mk_ap(db_session, number=f"AP-2026-05{i:02d}",
                     vendor=("ACME Ltd" if i < 2 else "Beta Corp"), src_id=uuid.uuid4())
    # give one a vendor invoice number to search on
    from app.models.ap_invoice import ApInvoice
    target = (await db_session.execute(select(ApInvoice).where(
        ApInvoice.ap_invoice_number == "AP-2026-0500"))).scalar_one()
    target.vendor_invoice_number = "INV-XYZ-77"
    target.nc_exported_at = datetime.now(timezone.utc)
    await db_session.flush()

    r = await client.get("/finance/v1/ap/invoices?limit=2&offset=0", headers=_h())
    body = r.json()
    assert body["total"] == 3 and len(body["items"]) == 2
    r2 = await client.get("/finance/v1/ap/invoices?limit=2&offset=2", headers=_h())
    assert len(r2.json()["items"]) == 1

    rq = await client.get("/finance/v1/ap/invoices?q=INV-XYZ", headers=_h())
    assert [i["ap_invoice_number"] for i in rq.json()["items"]] == ["AP-2026-0500"]
    rv = await client.get("/finance/v1/ap/invoices?q=beta", headers=_h())
    assert rv.json()["total"] == 1 and rv.json()["items"][0]["vendor_name"] == "Beta Corp"

    re_ = await client.get("/finance/v1/ap/invoices?exported=false", headers=_h())
    assert re_.json()["total"] == 2
    rt = await client.get("/finance/v1/ap/invoices?exported=true", headers=_h())
    assert rt.json()["total"] == 1 and rt.json()["items"][0]["nc_exported_at"] is not None


async def test_vendor_invoice_number_unique(client, db_session):
    vid = uuid.uuid4()
    from app.models.ap_invoice import ApInvoice
    a = await _mk_ap(db_session, number="AP-2026-0400", src_id=uuid.uuid4())
    a.vendor_id = vid; a.vendor_invoice_number = "DUP-1"
    await db_session.flush()
    b = ApInvoice(ap_invoice_number="AP-2026-0401", source="epms",
                  source_invoice_id=uuid.uuid4(), vendor_id=vid,
                  vendor_invoice_number="DUP-1", amount=Decimal("1"),
                  tax_amount=Decimal("0"), total_amount=Decimal("1"),
                  currency="CAD", invoice_date=date(2026, 7, 1), status="posted")
    db_session.add(b)
    from sqlalchemy.exc import IntegrityError
    with pytest.raises(IntegrityError):
        await db_session.flush()
    await db_session.rollback()
    # void rows are exempt — same number allowed again
    v = await _mk_ap(db_session, number="AP-2026-0402", src_id=uuid.uuid4())
    v.vendor_id = vid; v.vendor_invoice_number = "DUP-2"; v.status = "void"
    await db_session.flush()
    ok = await _mk_ap(db_session, number="AP-2026-0403", src_id=uuid.uuid4())
    ok.vendor_id = vid; ok.vendor_invoice_number = "DUP-2"
    await db_session.flush()          # must NOT raise
