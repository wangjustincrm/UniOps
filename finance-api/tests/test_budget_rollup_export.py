"""XLSX for the Budget-vs-Actual roll-up.

Pure function over a payload — no DB — so these run in milliseconds and pin the
things a spreadsheet can get wrong silently: missing rows, a "0.00" that should
have been blank, and a grouping that exports the other grouping's tree.
"""
import io

import openpyxl

from app.services.budget_rollup_export import build_rollup_xlsx

_METRICS = {
    "plan_period": "100.00", "actual_period": "60.00",
    "variance_period": "40.00", "variance_period_pct": "40.0",
    "plan_full_year": "120.00", "actual_ytd": "60.00",
    "remaining_full_year": "60.00", "consumed_pct": "50.0",
}


def _payload():
    leaf = {**_METRICS, "cost_center_id": "x", "cost_center_code": "GA-0107",
            "cost_center_name": "G&A-SC", "department_code": "0107",
            "department_name": "Supply Chain", "expense_centre": "GA"}
    node = {**_METRICS, "code": "GA", "label": "G&A", "children": [leaf]}
    dept = {**_METRICS, "code": "0107", "label": "Supply Chain", "children": [leaf]}
    return {
        "fiscal_year": 2026, "month_from": 1, "month_to": 9,
        "company": _METRICS,
        "by_expense_centre": [node],
        "by_department": [dept],
        "unallocated": {
            "actual_period": "900.00", "actual_ytd": "900.00",
            "breakdown": [{"key": "CRM09912", "label": "Shut-down loss",
                           "actual_period": "900.00", "actual_ytd": "900.00"}],
        },
        "reconciliation": {"placed_actual_period": "60.00",
                           "unallocated_actual_period": "900.00",
                           "total_actual_period": "960.00"},
    }


def _sheet(group_by="centre"):
    data = build_rollup_xlsx(rollup=_payload(), group_by=group_by, window_label="Jan-Sep")
    return openpyxl.load_workbook(io.BytesIO(data)).active


def _col_a(ws):
    return [ws.cell(r, 1).value for r in range(1, ws.max_row + 1)]


def test_every_cost_centre_is_written_whether_or_not_it_was_expanded():
    """The page collapses groups; the spreadsheet must not. A missing row in a
    file someone forwards is not a convenience, it is a wrong report."""
    labels = _col_a(_sheet())
    assert "G&A" in labels
    assert "G&A-SC" in labels          # the child, unexpanded on screen


def test_grouping_switches_the_tree():
    assert "Supply Chain" in _col_a(_sheet("department"))
    assert "Supply Chain" not in _col_a(_sheet("centre"))
    assert _sheet("department").cell(4, 1).value == "Department"
    assert _sheet("centre").cell(4, 1).value == "Expense centre"


def test_company_unallocated_and_reconciliation_all_present():
    labels = _col_a(_sheet())
    assert "Whole company" in labels
    assert "Category-level (not allocated)" in labels
    assert "Shut-down loss" in labels
    assert any(str(v).startswith("Total actual") for v in labels)
    # the proof line, so the file carries its own arithmetic
    assert any("in cost centres" in str(v) for v in labels)


def test_no_cell_is_mistaken_for_a_formula():
    """A text cell starting with "=" is a formula to Excel. The reconciliation
    line used to read "= 9,452,510 in cost centres + 15,711,858 category-level",
    which parses as nothing, so Excel dropped it and reported the workbook as
    damaged: "Removed Records: Formula from /xl/worksheets/sheet1.xml"."""
    ws = _sheet()
    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell.value, str):
                assert not cell.value.startswith("="), \
                    f"{cell.coordinate} looks like a formula: {cell.value!r}"


def test_absent_figures_stay_blank_rather_than_zero():
    """The unallocated rows have no plan. A 0.00 there reads as "budgeted
    nothing and spent it", which is a different statement from "not budgeted"."""
    ws = _sheet()
    row = next(r for r in range(1, ws.max_row + 1)
               if ws.cell(r, 1).value == "Shut-down loss")
    assert ws.cell(row, 3).value is None      # plan for the window
    assert ws.cell(row, 7).value is None      # full-year plan
    assert ws.cell(row, 4).value == 900.0     # actual is there


def test_header_is_frozen_and_money_is_formatted():
    ws = _sheet()
    assert ws.freeze_panes == "A5"
    row = next(r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value == "G&A")
    assert ws.cell(row, 3).number_format == "#,##0.00"
    assert ws.cell(row, 6).number_format == "0.0"     # percent column


def _styles_xml() -> str:
    import zipfile
    data = build_rollup_xlsx(rollup=_payload(), group_by="centre", window_label="Jan-Sep")
    return zipfile.ZipFile(io.BytesIO(data)).read("xl/styles.xml").decode()


def test_styles_excel_will_not_complain_about():
    """Excel opens a technically-valid xlsx and still says "we found a problem
    with some content" when the styles are nonsense to it — without naming the
    part. None of these show up in an openpyxl round-trip, which is exactly why
    they are asserted against the generated XML.

    All three were real: an `indent` with no `horizontal`, a solid fill missing
    its bgColor, and six-digit colours that become alpha-00 (transparent).
    """
    import re
    styles = _styles_xml()

    for tag in re.findall(r"<alignment[^>]*/>", styles):
        if "indent=" in tag:
            assert "horizontal=" in tag, f"indent without horizontal: {tag}"

    for fill in re.findall(r"<patternFill[^>]*>.*?</patternFill>", styles, re.S):
        if 'patternType="solid"' in fill:
            assert "<fgColor" in fill and "<bgColor" in fill, \
                f"solid fill missing a colour: {fill}"

    # openpyxl ships an indexed palette that is legitimately 00-prefixed; only
    # the colours this module sets are checked.
    ours = {"F5F5F5", "FAFAFA", "FFFBEB", "EDF2F7", "D4D4D4", "A3A3A3", "737373"}
    for rgb in re.findall(r'rgb="([0-9A-Fa-f]{8})"', styles):
        if rgb[2:].upper() in ours:
            assert rgb.upper().startswith("FF"), f"transparent colour: {rgb}"


def test_borders_carry_every_side():
    """A border element is expected to describe all of its sides; openpyxl
    happily writes one with only <top>."""
    import re
    for border in re.findall(r"<border>.*?</border>", _styles_xml(), re.S):
        for side in ("left", "right", "top", "bottom", "diagonal"):
            assert f"<{side}" in border, f"{side} missing from {border}"


def _monthly_payload():
    p = _payload()
    p["month_from"], p["month_to"] = 8, 9
    cells = [{"month": 8, "plan": "50.00", "actual": "20.00"},
             {"month": 9, "plan": "50.00", "actual": "40.00"}]
    for n in [p["company"], *p["by_expense_centre"], *p["by_department"],
              p["by_expense_centre"][0]["children"][0]]:
        n["monthly"] = cells
    p["company"] = {**_METRICS, "monthly": cells}
    p["unallocated"]["monthly"] = [{"month": 8, "actual": "900.00"},
                                   {"month": 9, "actual": "0.00"}]
    p["unallocated"]["breakdown"][0]["monthly"] = p["unallocated"]["monthly"]
    p["reconciliation"]["total_monthly"] = [{"month": 8, "actual": "920.00"},
                                            {"month": 9, "actual": "40.00"}]
    return p


def test_by_month_puts_plan_actual_pairs_before_the_window_columns():
    """Two columns a month — plan and actual, no monthly variance — ahead of the
    unchanged eight measures."""
    data = build_rollup_xlsx(rollup=_monthly_payload(), group_by="centre",
                             window_label="Aug-Sep", by_month=True)
    ws = openpyxl.load_workbook(io.BytesIO(data)).active
    head = [ws.cell(4, c).value for c in range(1, ws.max_column + 1)]
    assert head == ["Expense centre", "Code", "Aug Plan", "Aug Actual", "Sep Plan",
                    "Sep Actual", "Plan Aug-Sep", "Actual Aug-Sep", "Variance",
                    "Variance %", "Full-year plan", "Actual YTD", "Remaining", "Consumed %"]
    rows = {ws.cell(r, 1).value: [ws.cell(r, c).value for c in range(1, 15)]
            for r in range(5, ws.max_row + 1)}
    assert rows["Whole company"][2:8] == [50, 20, 50, 40, 100, 60]
    # category-level: no plan, blank rather than 0
    assert rows["Category-level (not allocated)"][2:6] == [None, 900, None, 0]
    # percentages still land on the percent format after the shift
    assert ws.cell(5, 10).number_format == "0.0"
    assert ws.cell(5, 14).number_format == "0.0"
    assert ws.cell(5, 3).number_format == "#,##0.00"


def test_without_by_month_the_layout_is_unchanged():
    ws = openpyxl.load_workbook(io.BytesIO(build_rollup_xlsx(
        rollup=_monthly_payload(), group_by="centre", window_label="Aug-Sep"))).active
    assert ws.max_column == 10
    assert ws.cell(4, 3).value == "Plan Aug-Sep"
