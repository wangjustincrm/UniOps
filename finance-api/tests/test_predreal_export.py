import io

import openpyxl

from app.services.predreal_export import build_budget_actual_xlsx


def _rows(data: bytes):
    wb = openpyxl.load_workbook(io.BytesIO(data))
    return list(wb.active.iter_rows(values_only=True))


def _find(rows, code):
    for r in rows:
        if r and r[1] == code:
            return r
    raise AssertionError(f"row not found: {code}")


def _sample():
    accounts = [
        {"account_id": "a1", "account_code": "5101010", "account_name": "Freight",
         "l1_code": "5101", "plan_by_month": {"1": "100.00", "2": "100.00"}, "plan_year": "1200.00"},
        {"account_id": "a2", "account_code": "6601010", "account_name": "Ads",
         "l1_code": "6601", "plan_by_month": {}, "plan_year": "0"},
    ]
    nc_monthly = {"a1": {1: "80.00", 2: "90.00"}}
    return accounts, nc_monthly


def _build():
    accounts, nc_monthly = _sample()
    return build_budget_actual_xlsx(accounts=accounts, nc_monthly=nc_monthly,
                                    fiscal_year=2026, cost_center_label="All Cost Centers")


def test_header_is_month_pairs():
    rows = _rows(_build())
    group = next(r for r in rows if r and r[0] == "Category")
    sub = rows[rows.index(group) + 1]
    # Labels occupy the first three columns; month names span their Plan/NC pair.
    assert group[:3] == ("Category", "Account Code", "Account Name")
    assert group[3] == "Jan" and group[5] == "Feb" and group[25] == "Dec" and group[27] == "Year"
    # merged cells leave the second column of each pair empty on the group row
    assert group[4] is None and group[26] is None and group[28] is None
    assert sub[3:9] == ("Plan", "NC", "Plan", "NC", "Plan", "NC")
    assert sub[27:29] == ("Plan", "NC")
    assert len(group) == 29


def test_account_row_carries_plan_and_nc_side_by_side():
    rows = _rows(_build())
    r = _find(rows, "5101010")
    assert r[0] == "5101" and r[2] == "Freight"
    assert r[3] == 100.0 and r[4] == 80.0      # Jan plan / Jan NC
    assert r[5] == 100.0 and r[6] == 90.0      # Feb plan / Feb NC
    assert r[27] == 1200.0 and r[28] == 170.0  # Year plan / Year NC


def test_one_row_per_account_no_vendor_breakdown():
    rows = _rows(_build())
    data_rows = [r for r in rows if r and r[1] in ("5101010", "6601010")]
    assert len(data_rows) == 2
    # no "(subtotal)" / vendor-name column survives anywhere in the sheet
    flat = {c for r in rows for c in r if isinstance(c, str)}
    assert "(subtotal)" not in flat and "Vendor" not in flat and "Metric" not in flat


def test_grand_total():
    rows = _rows(_build())
    gt = next(r for r in rows if r and r[0] == "Grand Total")
    assert gt[3] == 100.0 and gt[4] == 80.0
    assert gt[27] == 1200.0 and gt[28] == 170.0


# ── "Expand vendor": each account followed by its vendor breakdown ───────────

def _build_expanded():
    accounts, nc_monthly = _sample()
    partners = {"a1": [
        {"key": "v1", "partner_id": "v1", "partner_name": "Acme Freight", "source": "voucher",
         "by_month": {"1": "50.00", "2": "90.00"}, "year_total": "140.00"},
        {"key": "none", "partner_id": None, "partner_name": None, "source": "none",
         "by_month": {"1": "30.00"}, "year_total": "30.00"},
    ]}
    wb = openpyxl.load_workbook(io.BytesIO(build_budget_actual_xlsx(
        accounts=accounts, nc_monthly=nc_monthly, fiscal_year=2026,
        cost_center_label="All Cost Centers", partners=partners)))
    return wb.active


def test_expanded_header_adds_vendor_columns():
    rows = list(_build_expanded().iter_rows(values_only=True))
    group = next(r for r in rows if r and r[0] == "Category")
    assert group[:5] == ("Category", "Account Code", "Account Name", "Vendor", "Vendor Source")
    assert group[5] == "Jan" and group[29] == "Year" and len(group) == 31


def test_expanded_account_followed_by_its_vendors():
    ws = _build_expanded()
    rows = list(ws.iter_rows(values_only=True))
    i = next(k for k, r in enumerate(rows) if r and r[1] == "5101010" and r[3] in (None, ""))
    acct, v1, v2 = rows[i], rows[i + 1], rows[i + 2]
    assert acct[5] == 100.0 and acct[6] == 80.0 and acct[30] == 170.0
    # vendor rows leave the account labels empty: the account reads once above them
    assert v1[3:5] == ("Acme Freight", "from voucher")
    assert all(c in (None, "") for c in v1[:3]) and all(c in (None, "") for c in v2[:3])
    assert v1[5] is None and v1[6] == 50.0 and v1[8] == 90.0   # no plan per vendor
    assert v1[29] is None and v1[30] == 140.0
    assert v2[3] == "(no vendor — internal entry)" and v2[4] in (None, "")
    assert v2[6] == 30.0 and v2[30] == 30.0
    # collapsible under the account row
    assert ws.row_dimensions[i + 2].outline_level == 1
    assert ws.row_dimensions[i + 1].outline_level == 0
    # an account with no NC actual has no vendor rows
    k = next(k for k, r in enumerate(rows) if r and r[1] == "6601010")
    assert k + 1 == len(rows) or rows[k + 1][3] in (None, "")


def test_expanded_grand_total_counts_accounts_only():
    rows = list(_build_expanded().iter_rows(values_only=True))
    gt = next(r for r in rows if r and r[0] == "Grand Total")
    assert gt[5] == 100.0 and gt[6] == 80.0
    assert gt[29] == 1200.0 and gt[30] == 170.0
